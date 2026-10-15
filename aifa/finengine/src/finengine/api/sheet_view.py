"""Чтение одного листа xlsx-источника для отрисовки сеткой.

Нужно рабочей области: Excel-артефакт (метрики_*.xlsx, factor_analysis_*.xlsx)
показывается таблицей прямо в UI, а не только скачивается. Существующая выдача
/sources/{id}/text для этого не годится — она склеивает ячейки в TSV и
выбрасывает пустые строки (`if any(cells)`), из-за чего адреса ячеек съезжают,
а формулы теряются вовсе.

Отдаём РАЗРЕЖЕННЫЙ список ячеек, а не матрицу строк: у финансовых выгрузок
пустых ячеек больше, чем заполненных, и плотная матрица — это в основном
пересылка null'ов. Адреса r/c 1-based, как в самом Excel.

Два прохода openpyxl сделаны намеренно: `data_only=True` отдаёт вычисленное
значение, `data_only=False` — исходную формулу; за один проход обе величины
получить нельзя. Второй проход не роскошь: xlsx, собранный самим openpyxl,
кэша значений НЕ содержит, и формульные ячейки в первом проходе приходят как
None — без формулы UI показал бы пустоту вместо содержимого.
"""
from __future__ import annotations

import datetime as dt
import io
import math
from decimal import Decimal
from typing import Any

# Капы выдачи. Без них один «широкий» лист (max_row/max_column у Excel
# раздуваются от форматирования пустых ячеек) укладывает и сервер, и вкладку.
MAX_ROWS = 2_000
MAX_COLS = 200
# Потолок на площадь одного ответа: 2000×200 — это уже мегабайты JSON.
MAX_TOTAL_CELLS = 100_000
# openpyxl держит разобранную книгу в памяти целиком, причём дважды
# (два прохода), поэтому крупные файлы сеткой не отдаём вовсе.
MAX_FILE_BYTES = 15 * 1024 * 1024

# Форматы, которые openpyxl действительно читает. Легаси .xls (BIFF) —
# не читает, и это надо сказать вслух, а не отдавать пустую сетку.
GRID_EXTENSIONS = {".xlsx", ".xlsm"}

# Значения ошибок Excel приезжают из кэша обычными строками.
_ERROR_LITERALS = frozenset({
    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#GETTING_DATA",
})


class SheetReadError(Exception):
    """Файл не разбирается как xlsx — вызывающий превращает в 4xx."""


class SheetNotFound(Exception):
    """Запрошенного листа в книге нет — вызывающий превращает в 404."""


def read_sheet(
    data: bytes,
    sheet: str = "",
    max_rows: int = MAX_ROWS,
    max_cols: int = MAX_COLS,
) -> dict[str, Any]:
    """Разобрать байты xlsx и вернуть один лист сеткой.

    `sheet` пустой — берётся первый лист книги. В ответе `sheets` — имена
    всех листов (вкладки), `n_rows`/`n_cols` — настоящий размер листа:
    `cells` может покрывать меньше из-за капов, и тогда `truncated` говорит
    об этом явно, чтобы UI не выдавал обрезок за целое.
    """
    import openpyxl

    max_rows = max(1, min(max_rows, MAX_ROWS))
    max_cols = max(1, min(max_cols, MAX_COLS))

    try:
        wb_values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        wb_formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    except Exception as e:  # noqa: BLE001 — битый/чужой формат, наверх уходит 4xx
        raise SheetReadError(f"{e.__class__.__name__}: {e}") from e

    try:
        names = list(wb_values.sheetnames)
        if not names:
            raise SheetReadError("в книге нет листов")
        if sheet and sheet not in names:
            raise SheetNotFound(sheet)
        name = sheet or names[0]
        out = _read_one(wb_values[name], wb_formulas[name], max_rows, max_cols)
        out["sheets"] = names
        out["sheet"] = name
        return out
    finally:
        wb_values.close()
        wb_formulas.close()


def _read_one(ws_v, ws_f, max_rows: int, max_cols: int) -> dict[str, Any]:
    # Лист диаграммы (Chartsheet) ячеек не имеет — отдаём пустым, а не падаем.
    if not hasattr(ws_v, "iter_rows"):
        return {"n_rows": 0, "n_cols": 0, "truncated": False, "cells": []}

    n_rows = ws_v.max_row or 0
    n_cols = ws_v.max_column or 0
    row_limit = min(n_rows, max_rows)
    col_limit = min(n_cols, max_cols)
    if row_limit and col_limit:
        # Площадь режем по строкам: колонки нужны все, иначе таблица теряет
        # правые периоды и выглядит достоверно-неверной.
        row_limit = min(row_limit, MAX_TOTAL_CELLS // col_limit)

    cells: list[dict[str, Any]] = []
    if row_limit and col_limit:
        value_rows = ws_v.iter_rows(min_row=1, max_row=row_limit, min_col=1, max_col=col_limit)
        formula_rows = ws_f.iter_rows(min_row=1, max_row=row_limit, min_col=1, max_col=col_limit)
        for row_v, row_f in zip(value_rows, formula_rows):
            for cell_v, cell_f in zip(row_v, row_f):
                value = _json_value(cell_v.value)
                formula = _formula_of(cell_f.value)
                if value is None and formula is None:
                    continue  # разреженность: пустые ячейки не шлём
                item: dict[str, Any] = {
                    "r": cell_v.row,
                    "c": cell_v.column,
                    "v": value,
                    "t": _cell_type(value),
                }
                if formula is not None:
                    item["f"] = formula
                font = cell_v.font
                if font is not None and font.bold:
                    item["b"] = True
                cells.append(item)

    out: dict[str, Any] = {
        "n_rows": n_rows,
        "n_cols": n_cols,
        "truncated": row_limit < n_rows or col_limit < n_cols,
        "cells": cells,
    }
    widths = _col_widths(ws_v, col_limit)
    if widths:
        out["col_widths"] = widths
    return out


def _col_widths(ws, col_limit: int) -> list[float]:
    """Ширины колонок в символах Excel. 0 — ширина не задана, UI берёт свою."""
    from openpyxl.utils import get_column_letter

    dims = ws.column_dimensions
    widths: list[float] = []
    has_any = False
    for c in range(1, col_limit + 1):
        # .get(), а не [], — обращение по ключу создаёт запись с дефолтами.
        dim = dims.get(get_column_letter(c))
        width = float(dim.width) if dim is not None and dim.width else 0.0
        has_any = has_any or width > 0
        widths.append(round(width, 1))
    return widths if has_any else []


def _json_value(v: Any) -> Any:
    """Значение ячейки → JSON-совместимый скаляр."""
    if v is None or isinstance(v, (str, bool, int)):
        return v
    if isinstance(v, float):
        # NaN/Infinity ломают JSON.parse на фронте — честнее отдать пусто.
        return v if math.isfinite(v) else None
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, dt.datetime):
        # openpyxl расширяет дату до datetime, и в финансовой таблице каждая
        # дата отчёта приезжала бы как «2024-12-31T00:00:00». Полночь без
        # number_format неотличима от чистой даты — отдаём датой.
        return v.date().isoformat() if v.time() == dt.time.min else v.isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return v.total_seconds()
    return str(v)


def _cell_type(v: Any) -> str:
    """s|n|d|b|e — UI выравнивает по нему содержимое ячейки."""
    if isinstance(v, bool):
        return "b"          # проверять до int: bool — его подкласс
    if isinstance(v, (int, float)):
        return "n"
    if isinstance(v, str):
        if v in _ERROR_LITERALS:
            return "e"
        # Даты и время отданы ISO-строкой; помечаем их как даты, иначе UI
        # выровняет отчётную дату как текст.
        return "d" if _looks_like_iso_date(v) else "s"
    return "s"


def _looks_like_iso_date(v: str) -> bool:
    if not (10 <= len(v) <= 26 and v[:4].isdigit() and v[4:5] == "-"):
        return False
    try:
        dt.datetime.fromisoformat(v)
    except ValueError:
        return False
    return True


def _formula_of(v: Any) -> str | None:
    """Формула ячейки строкой, если это формула."""
    if isinstance(v, str):
        return v if v.startswith("=") else None
    # ArrayFormula / DataTableFormula — объекты openpyxl с текстом формулы.
    text = getattr(v, "text", None)
    return text if isinstance(text, str) and text.startswith("=") else None

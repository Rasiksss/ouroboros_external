from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any

from finengine.roles.access import ALL_EXTRACTION_ROLE
from finengine.roles.scope import RolePlan
from finengine.skills.periods import normalize_incomplete_duration_label, normalize_period_label


MISSING_TEXT = "Не найдено"

# Код реестра внутри текста: то же выражение, что у `formulaRu` на фронте.
_REGISTRY_CODE_RE = re.compile(r"[A-Z][A-Z0-9_]{2,}")

_DIE_SEGMENT_TITLES: dict[str, str] = {
    "SALES_REVENUE": "Структура выручки",
    "GROSS_EXPENSES": "Структура валовой прибыли",
    "OPERATING_PROFIT": "Структура операционной прибыли",
    "NET_INCOME": "Структура чистой прибыли",
    "COSTS_EXPENSES": "Структура себестоимости",
}


@dataclass(frozen=True)
class PeriodInfo:
    raw: str
    kind: str
    year: int
    sort_key: tuple[int, int, int, str]
    family: str


@dataclass(frozen=True)
class PeriodCompactContext:
    annual_years: frozenset[int]
    quarter_keys: frozenset[str]


_QUARTER_ENDS: dict[tuple[int, int], str] = {
    (3, 31): "Q1",
    (6, 30): "Q2",
    (9, 30): "Q3",
    (12, 31): "Q4",
}


def _parse_period(period: object) -> PeriodInfo | None:
    raw = str(period or "").strip()
    if not raw:
        return None
    if re.fullmatch(r"\d{4}", raw):
        year = int(raw)
        return PeriodInfo(raw, "fy", year, (year, 12, 31, raw), "FY")
    try:
        dt = datetime.strptime(raw, "%d.%m.%Y")
        return PeriodInfo(raw, "date", dt.year, (dt.year, dt.month, dt.day, raw), "DATE")
    except ValueError:
        pass

    upper = raw.upper()
    if m := re.fullmatch(r"FY\s+(\d{4})", upper):
        year = int(m.group(1))
        return PeriodInfo(raw, "fy", year, (year, 12, 30, raw), "FY")
    if m := re.fullmatch(r"Q([1-4])\s+(\d{4})", upper):
        quarter = int(m.group(1))
        year = int(m.group(2))
        return PeriodInfo(raw, "quarter", year, (year, quarter * 3, 20, raw), f"Q{quarter}")
    if m := re.fullmatch(r"(\d{1,2})M\s+(\d{4})", upper):
        month = int(m.group(1))
        year = int(m.group(2))
        if 1 <= month <= 12:
            return PeriodInfo(raw, "months", year, (year, month, 10, raw), f"{month}M")
    if m := re.fullmatch(r"(\d{1,2})M\s+(\d{1,2}\.\d{1,2}\.\d{4})", upper):
        month = int(m.group(1))
        try:
            dt = datetime.strptime(m.group(2), "%d.%m.%Y")
        except ValueError:
            return None
        return PeriodInfo(raw, "rolling_months", dt.year, (dt.year, dt.month, dt.day, raw), f"{month}M")
    if m := re.fullmatch(r"(\d{1,2})-(\d{1,2})M\s+(\d{4})", upper):
        start_month = int(m.group(1))
        end_month = int(m.group(2))
        year = int(m.group(3))
        if 1 <= start_month <= end_month <= 12:
            return PeriodInfo(raw, "month_range", year, (year, end_month, 15, raw), f"{start_month}-{end_month}M")
    return None


def _duration_group(info: PeriodInfo | None) -> int:
    if not info:
        return 9
    if info.kind == "date":
        return 0
    months = _duration_months(info)
    if info.kind == "quarter" or months == 3 or _month_range_length(info) == 3:
        return 1
    if months == 6 or _month_range_length(info) == 6:
        return 2
    if months == 9 or _month_range_length(info) == 9:
        return 3
    if info.kind == "fy" or months == 12 or _month_range_length(info) == 12:
        return 4
    return 8


def _period_sort_key(period: str) -> tuple[int, int, int, int, str]:
    info = _parse_period(period)
    if not info:
        return (9, 9999, 99, 99, str(period))
    year, month, day, raw = info.sort_key
    return (_duration_group(info), year, month, day, raw)


def _quarter_display_key(period: object, context: PeriodCompactContext | None = None) -> str | None:
    info = _parse_period(period)
    if not info:
        return None
    if info.kind == "quarter":
        return f"{info.family} {info.year}"
    if info.kind != "date":
        return None
    try:
        dt = datetime.strptime(str(period or "").strip(), "%d.%m.%Y")
    except ValueError:
        return None
    quarter = _QUARTER_ENDS.get((dt.month, dt.day))
    if not quarter:
        return None
    key = f"{quarter} {dt.year}"
    if quarter == "Q4" and context and dt.year in context.annual_years:
        return None
    if context is None or key not in context.quarter_keys:
        return None
    return key


def _annual_display_key(
    period: object,
    context: PeriodCompactContext | None = None,
) -> str | None:
    info = _parse_period(period)
    if not info:
        return None
    if info.kind == "fy":
        return str(info.year)
    if info.kind == "months" and info.family == "12M":
        return str(info.year)
    if info.kind == "rolling_months" and info.family == "12M":
        raw = str(period or "").strip()
        if re.search(r"\b31\.12\.(20\d{2})\b", raw):
            return str(info.year)
    if info.kind == "date":
        try:
            dt = datetime.strptime(str(period or "").strip(), "%d.%m.%Y")
        except ValueError:
            return None
        if dt.day == 31 and dt.month == 12:
            if context and f"Q4 {dt.year}" in context.quarter_keys and dt.year not in context.annual_years:
                return None
            return str(dt.year)
    return None


def _period_compact_context(periods: list[str]) -> PeriodCompactContext:
    annual_years: set[int] = set()
    quarter_keys: set[str] = set()
    for period in periods:
        info = _parse_period(period)
        if not info:
            continue
        if info.kind == "fy":
            annual_years.add(info.year)
        elif info.kind == "months" and info.family == "12M":
            annual_years.add(info.year)
        elif info.kind == "rolling_months" and info.family == "12M":
            raw = str(period or "").strip()
            if re.search(r"\b31\.12\.(20\d{2})\b", raw):
                annual_years.add(info.year)
        elif info.kind == "quarter":
            quarter_keys.add(f"{info.family} {info.year}")
    return PeriodCompactContext(frozenset(annual_years), frozenset(quarter_keys))


def _meta_period_context(meta: dict[str, Any]) -> PeriodCompactContext | None:
    context = meta.get("_period_compact_context")
    return context if isinstance(context, PeriodCompactContext) else None


def _with_period_context(metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    if _meta_period_context(meta):
        return meta
    out = dict(meta)
    out["_period_compact_context"] = _period_compact_context(_raw_periods(metrics, meta))
    return out


def _compact_period_key(
    period: object,
    context: PeriodCompactContext | None = None,
) -> str:
    raw = str(period or "").strip()
    return _annual_display_key(raw, context) or _quarter_display_key(raw, context) or raw


def _compact_periods(
    periods: list[str],
    context: PeriodCompactContext | None = None,
) -> list[str]:
    context = context or _period_compact_context(periods)
    keys = {
        key
        for period in periods
        if (key := _compact_period_key(period, context))
    }
    return sorted(keys, key=_period_sort_key)


def _period_anchor_labels(metric: dict[str, Any] | None, meta: dict[str, Any]) -> list[object]:
    anchors: list[object] = []
    anchors.extend(meta.get("periods") or [])
    anchors.extend([meta.get("period_current"), meta.get("period_prev")])
    for pv in (metric or {}).get("period_values") or []:
        if isinstance(pv, dict):
            anchors.append(pv.get("period"))
            anchors.append(pv.get("raw_period"))
    return [anchor for anchor in anchors if str(anchor or "").strip()]


def _normalize_period_key(
    period: object,
    metric: dict[str, Any] | None,
    meta: dict[str, Any],
) -> str:
    raw = str(period or "").strip()
    if not raw or _parse_period(raw):
        return raw
    normalized = normalize_period_label(raw)
    if normalized and _parse_period(normalized):
        return normalized
    incomplete = normalize_incomplete_duration_label(
        raw,
        tuple(_period_anchor_labels(metric, meta)),
    )
    return incomplete or raw


def _calendar_quarter_key(dt: datetime) -> str:
    candidates = [
        (dt.year - 1, 4, datetime(dt.year - 1, 12, 31)),
        (dt.year, 1, datetime(dt.year, 3, 31)),
        (dt.year, 2, datetime(dt.year, 6, 30)),
        (dt.year, 3, datetime(dt.year, 9, 30)),
        (dt.year, 4, datetime(dt.year, 12, 31)),
        (dt.year + 1, 1, datetime(dt.year + 1, 3, 31)),
    ]
    year, quarter, _ = min(
        candidates,
        key=lambda item: (abs((item[2] - dt).days), 0 if item[2] >= dt else 1),
    )
    return f"Q{quarter} {year}"


def _duration_months(info: PeriodInfo) -> int | None:
    if info.kind not in {"months", "rolling_months"}:
        return None
    try:
        return int(info.family.replace("M", ""))
    except ValueError:
        return None


def _display_period(period: object, *, style: str = "default") -> str:
    raw = str(period or "").strip()
    annual = _annual_display_key(raw)
    if annual:
        return annual
    info = _parse_period(raw)
    if not info:
        return raw
    if style == "die":
        if info.kind == "quarter":
            quarter = info.family.replace("Q", "")
            return f"{quarter}Q{str(info.year)[-2:]}"
        if info.kind == "fy":
            return str(info.year)
        return raw
    if info.kind == "fy":
        return str(info.year)
    if info.kind == "quarter":
        return f"{info.family.replace('Q', '')} квартал {info.year}"
    if info.kind == "months":
        months = info.family.replace("M", "")
        return f"за {months} месяцев {info.year}"
    if info.kind == "month_range":
        start, end = info.family.replace("M", "").split("-", 1)
        return f"за {start}-{end} месяцев {info.year}"
    if info.kind == "rolling_months":
        months = info.family.replace("M", "")
        return f"за {months} месяцев до {raw.split(maxsplit=1)[1]}"
    return raw


def _legacy_period_label(period: object) -> str:
    """Old formatter convention: FY becomes a year, quarters/YTD stay raw."""
    raw = str(period or "").strip()
    annual = _annual_display_key(raw)
    if annual:
        return annual
    info = _parse_period(raw)
    if info and info.kind == "fy":
        return str(info.year)
    return raw


def _die_period_key(
    period: object,
    *,
    mode: str = "quarter",
    aliases: dict[str, str] | None = None,
) -> str | None:
    raw = str(period or "").strip()
    info = _parse_period(raw)
    if not info:
        return None
    months = _duration_months(info)
    key: str | None = None
    if mode == "annual":
        annual = _annual_display_key(raw)
        if annual:
            key = f"FY {annual}"
        elif months == 12:
            key = raw
        elif info.kind == "date":
            key = raw
        elif info.kind == "quarter":
            key = f"{info.family} {info.year}"
        elif info.kind in {"months", "rolling_months", "month_range"}:
            key = raw
    elif info.kind == "quarter":
        key = f"{info.family} {info.year}"
    elif info.kind == "fy":
        key = f"FY {info.year}"
    elif info.kind == "date":
        key = raw
    elif info.kind in {"months", "rolling_months", "month_range"}:
        key = raw
    if key and aliases:
        return aliases.get(key, key)
    return key


def _die_period_priority(period: object, *, mode: str = "quarter") -> int:
    info = _parse_period(period)
    if not info:
        return 0
    if info.kind == "rolling_months":
        return 80
    if info.kind == "month_range":
        return 70
    if info.kind == "months":
        return 65
    if info.kind in {"quarter", "fy"}:
        return 55
    if info.kind == "date":
        return 45 if mode == "annual" else 35
    return 10


def _rolling_period_date(period: str) -> datetime | None:
    parts = str(period or "").strip().split(maxsplit=1)
    if len(parts) != 2:
        return None
    try:
        return datetime.strptime(parts[1], "%d.%m.%Y")
    except ValueError:
        return None


def _die_period_aliases(
    metrics: list[dict[str, Any]],
    meta: dict[str, Any],
    *,
    mode: str,
) -> dict[str, str]:
    aliases: dict[str, str] = {}
    for period in _raw_periods(metrics, meta):
        info = _parse_period(period)
        if not info or info.kind != "rolling_months":
            continue
        months = _duration_months(info)
        if not months:
            continue
        if months == 12:
            aliases[f"FY {info.year}"] = info.raw
            aliases[str(info.year)] = info.raw
        elif months:
            aliases[f"{months}M {info.year}"] = info.raw
            dt = _rolling_period_date(info.raw)
            if months == 3 and dt:
                aliases[_calendar_quarter_key(dt)] = info.raw
    return aliases


def _raw_periods(metrics: list[dict[str, Any]], meta: dict[str, Any]) -> list[str]:
    periods: list[str] = []
    for period in meta.get("periods") or []:
        raw = str(period or "").strip()
        if raw:
            periods.append(raw)
    for period in (meta.get("period_prev"), meta.get("period_current")):
        raw = str(period or "").strip()
        if raw:
            periods.append(raw)
    for metric in metrics:
        for pv in metric.get("period_values") or []:
            if not isinstance(pv, dict):
                continue
            raw = str(pv.get("period") or "").strip()
            if raw:
                periods.append(raw)
    return periods


def _is_annual_die_period(period: object) -> bool:
    info = _parse_period(period)
    if not info:
        return False
    if info.kind == "fy":
        return True
    if info.kind in {"months", "rolling_months"} and info.family == "12M":
        return True
    return False


def _is_quarterly_die_period(period: object) -> bool:
    info = _parse_period(period)
    if not info:
        return False
    if info.kind == "quarter":
        return True
    return info.kind in {"months", "rolling_months", "month_range"} and not _is_annual_die_period(period)


def _month_range_length(info: PeriodInfo) -> int | None:
    if info.kind != "month_range":
        return None
    try:
        start, end = (int(part) for part in info.family.replace("M", "").split("-", 1))
    except ValueError:
        return None
    return end - start + 1


def _is_die_output_period(period: object) -> bool:
    info = _parse_period(period)
    if not info:
        return False
    if info.kind in {"date", "fy", "quarter"}:
        return True
    months = _duration_months(info)
    if months in {3, 12}:
        return True
    return _month_range_length(info) in {3, 12}


def _die_period_mode(metrics: list[dict[str, Any]], meta: dict[str, Any]) -> str:
    periods = _raw_periods(metrics, meta)
    if any(_is_annual_die_period(period) for period in periods):
        return "annual"
    if any(_is_quarterly_die_period(period) for period in periods):
        return "quarter"
    year_end_dates = [
        period
        for period in periods
        if (info := _parse_period(period)) and info.kind == "date" and _annual_display_key(period)
    ]
    return "annual" if year_end_dates else "quarter"


def _value_is_present(value: Any) -> bool:
    return value is not None and value != "" and value != MISSING_TEXT


def _to_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text == MISSING_TEXT:
        return None
    text = text.replace("\u00a0", " ").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _unit_currency(unit: str) -> tuple[str, str]:
    clean = (unit or "").replace("\u00a0", " ").replace(".", "").strip()
    clean = re.sub(r"\s+", " ", clean)
    if not clean:
        return "", ""
    first, _, rest = clean.partition(" ")
    if first in {"%", "раз"}:
        return first, ""
    return first, rest


def _metric_has_values(metric: dict[str, Any] | None) -> bool:
    if not metric:
        return False
    if _value_is_present(metric.get("value_current")) or _value_is_present(metric.get("value_prev")):
        return True
    for pv in metric.get("period_values") or []:
        if isinstance(pv, dict) and _value_is_present(pv.get("value")):
            return True
    return False


def _metric_index(metrics: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    by_enum: dict[str, list[dict[str, Any]]] = {}
    for metric in metrics:
        enum = str(metric.get("enum") or "").strip().upper()
        if not enum:
            continue
        by_enum.setdefault(enum, []).append(metric)

    chosen: dict[str, dict[str, Any]] = {}
    for enum, rows in by_enum.items():
        rows = sorted(
            rows,
            key=lambda row: (
                0 if _metric_has_values(row) else 1,
                0 if not str(row.get("segment") or "").strip()
                else 1 if str(row.get("segment") or "").strip().upper() == "TOTAL"
                else 2,
                -len([pv for pv in row.get("period_values") or [] if isinstance(pv, dict) and _value_is_present(pv.get("value"))]),
            ),
        )
        chosen[enum] = rows[0]
    return chosen


def _period_values(
    metric: dict[str, Any] | None,
    meta: dict[str, Any],
    *,
    die: bool = False,
    die_period_mode: str = "quarter",
    die_period_aliases: dict[str, str] | None = None,
) -> dict[str, dict[str, Any]]:
    if not metric:
        return {}
    out: dict[str, dict[str, Any]] = {}
    priorities: dict[str, int] = {}
    period_context = _meta_period_context(meta)

    def put(period: str, item: dict[str, Any], priority: int = 0) -> None:
        current_priority = priorities.get(period, -1)
        if priority >= current_priority:
            out[period] = item
            priorities[period] = priority

    for pv in metric.get("period_values") or []:
        if not isinstance(pv, dict) or not _value_is_present(pv.get("value")):
            continue
        period = _normalize_period_key(pv.get("period"), metric, meta)
        priority = 0
        if die:
            priority = _die_period_priority(period, mode=die_period_mode)
            period = _die_period_key(period, mode=die_period_mode, aliases=die_period_aliases) or ""
        if not period:
            continue
        item = {
            "value": pv.get("value"),
            "unit": str(pv.get("unit") or metric.get("unit") or meta.get("unit") or ""),
            "page": pv.get("page") or metric.get("page"),
            "source_form": pv.get("source_form") or metric.get("source_form") or "",
            "source_text": pv.get("source_text") or "",
        }
        put(period, item, priority)
        compact_key = _compact_period_key(period, period_context)
        if compact_key:
            put(compact_key, item, priority)

    # ⚠️ ЗАПАСНОЙ ПУТЬ ДОПОЛНЯЕТ, А НЕ ЗАМЕЩАЕТ (09.08.2026, дефект Русгидро).
    #
    # Здесь стояло `if not out:` — то есть `value_current`/`value_prev`
    # смотрели ТОЛЬКО когда `period_values` пуст целиком. Движок же заполняет
    # `period_values` ЧАСТИЧНО: у показателя есть запись за текущий год и нет
    # за предыдущий, при том что `value_prev` заполнен. Итог на экране
    # владельца: колонка FY 2023 есть (её создали те немногие показатели, у
    # которых запись за 2023 была), а значений в ней нет почти нигде — при том
    # что markdown-таблица в чате рисуется ПО ДРУГИМ полям и оба периода
    # показывает. Две раскладки одного прогона расходились молча.
    #
    # Теперь дополняем по каждому периоду отдельно и ТОЛЬКО там, где значения
    # ещё нет: `period_values` остаётся главным источником, и его число никогда
    # не перетирается менее точным.
    cur = meta.get("period_current") or "Текущий"
    prev = meta.get("period_prev") or "Предыдущий"
    fallback = ((cur, metric.get("value_current")), (prev, metric.get("value_prev")))
    for period, value in fallback:
        if not _value_is_present(value):
            continue
        period = _normalize_period_key(period, metric, meta)
        key = (
            _die_period_key(period, mode=die_period_mode, aliases=die_period_aliases)
            if die
            else str(period)
        )
        if not key or key in out:
            continue
        item = {
            "value": value,
            "unit": metric.get("unit") or meta.get("unit") or "",
            "page": metric.get("page"),
            "source_form": metric.get("source_form") or "",
        }
        priority = _die_period_priority(period, mode=die_period_mode) if die else 0
        put(key, item, priority)
        compact_key = _compact_period_key(key, period_context)
        if compact_key and compact_key not in out:
            put(compact_key, item, priority)
    return out


def _clean_source_form(value: object) -> str:
    return str(value or "").strip()


def _company_name_from_meta(meta: dict[str, Any]) -> str:
    for key in ("company_name", "bank_name", "company"):
        value = str(meta.get(key) or "").strip()
        if value:
            return value
    return ""


def _format_source_ref(page: int, source_form: object = "") -> str:
    if page <= 0:
        return ""
    form = _clean_source_form(source_form)
    return f"стр. {page} ({form})" if form else f"стр. {page}"


def _add_source_ref(refs: dict[int, str], page: object, source_form: object = "") -> None:
    try:
        page_int = int(page or 0)
    except Exception:
        page_int = 0
    if page_int <= 0:
        return
    form = _clean_source_form(source_form)
    if page_int not in refs or (form and not refs[page_int]):
        refs[page_int] = form


def _all_periods(
    metrics: list[dict[str, Any]],
    meta: dict[str, Any],
    *,
    die: bool = False,
    die_period_mode: str = "quarter",
    die_period_aliases: dict[str, str] | None = None,
) -> list[str]:
    meta = _with_period_context(metrics, meta)
    period_context = _meta_period_context(meta)
    periods: set[str] = set()
    for metric in metrics:
        periods.update(
            _period_values(
                metric,
                meta,
                die=die,
                die_period_mode=die_period_mode,
                die_period_aliases=die_period_aliases,
            ).keys()
        )
    if not periods:
        for p in (meta.get("period_prev"), meta.get("period_current")):
            key = (
                _die_period_key(p, mode=die_period_mode, aliases=die_period_aliases)
                if die
                else str(p or "").strip()
            )
            if key:
                periods.add(key)
    return _compact_periods(list(periods), period_context)


def _source(metric: dict[str, Any] | None) -> str:
    if not metric:
        return ""
    refs: dict[int, str] = {}
    _add_source_ref(refs, metric.get("page"), metric.get("source_form"))
    for pv in metric.get("period_values") or []:
        if not isinstance(pv, dict):
            continue
        _add_source_ref(refs, pv.get("page"), pv.get("source_form") or metric.get("source_form"))
    return ", ".join(_format_source_ref(page, refs.get(page)) for page in sorted(refs))


def _in_russian(text: str, names: dict[str, str]) -> str:
    """Коды реестра внутри любого текста — русскими именами из той же книги."""
    if not names:
        return text
    return _REGISTRY_CODE_RE.sub(lambda m: names.get(m.group(0)) or m.group(0), text)


def _ru_number(value: object) -> str:
    """Число так же, как его печатает таблица на экране.

    Экран зовёт `toLocaleString("ru-RU")`: разряды неразрывным пробелом, дробь
    запятой, не больше двух знаков, у целого дроби нет. Округление —
    «от нуля» на ровной половине, поэтому Decimal с ROUND_HALF_UP, а не
    `format`: тот на 1000,125 даёт 1000,12, и книга разошлась бы с экраном на
    копейку — ровно тот класс расхождения, ради которого правило свели в одно.
    """
    if value is None:
        return "—"
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    rounded = Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    whole, _, fraction = f"{abs(rounded):f}".partition(".")
    fraction = fraction.rstrip("0")
    grouped = f"{int(whole):,}".replace(",", " ")
    sign = "-" if rounded < 0 else ""
    return f"{sign}{grouped}" + (f",{fraction}" if fraction else "")


def _calculated_source(metric: dict[str, Any], names: dict[str, str]) -> str:
    """«Источник» расчётной строки — её формула с подставленными числами.

    ЦЕНА (Б5, 19.08.2026). В ролевых книгах слова `formula` не было ни разу на
    1787 строк: у расчётной строки колонка «Источник» оставалась ПУСТОЙ, потому
    что `_source` собирает ссылки по номеру страницы, а у расчёта страницы нет.
    Аналитик получал в книге число без единого следа происхождения.

    Отдельной колонки не заводим: у расчётной строки это место и так свободно, а
    источник числа, посчитанного формулой, — и есть его формула. Русские имена
    входов берём из соседних строк той же книги — второго словаря
    «код → имя» не заводим, он и так приезжает вместе с числами.

    ЦЕНА ВТОРАЯ (20.08.2026, состязательная проверка). Первая редакция печатала
    здесь сырые коды реестра и сырой float (`OPERATING_PROFIT`, `28355.0`), а
    экран — переведённое и отформатированное (`Операционная прибыль`,
    `28 355`). Аналитик, сверявший книгу с экраном, получал два разных ответа
    на один вопрос «как это посчитано». Правило теперь ОДНО и закреплено
    общей разметкой `tests/fixtures/calculation_signature.json`, которую
    читают оба сторожа — pytest и vitest.
    """
    parts: list[str] = []
    formula = str(metric.get("formula") or "").strip()
    if formula:
        parts.append(_in_russian(formula, names))
    items: list[str] = []
    periods: set[str] = set()
    for item in metric.get("inputs") or []:
        label = _in_russian(str(item.get("enum") or "").strip(), names)
        transform = str(item.get("transform") or "")
        if transform == "avg":
            label = f"AVG({label})"
        elif transform == "abs":
            label = f"abs({label})"
        elif transform == "sign":
            # Знак отчёта у расчётного оттока снимается ОТРИЦАНИЕМ, и назвать
            # его модулем нельзя: у входящего нетто-потока `abs(...)` не даёт
            # напечатанного числа (`_formula_inputs`).
            label = f"-{label}"
        zero = " — подставлен ноль" if item.get("substituted_zero") else ""
        items.append(f"{label} = {_ru_number(item.get('value'))}{zero}")
        periods.add(str(item.get("period") or "").strip())
    if items:
        period = periods.pop() if len(periods) == 1 else ""
        head = f"входы за {period}: " if period else "входы: "
        parts.append(head + "; ".join(items))
    missing = [_in_russian(str(enum), names) for enum in (metric.get("calc_missing") or [])]
    if missing:
        parts.append("не хватило: " + ", ".join(missing))
    return " · ".join(parts)


def _source_from_value_item(item: dict[str, Any], metric: dict[str, Any] | None) -> str:
    page = item.get("page") or (metric or {}).get("page")
    source_form = item.get("source_form") or (metric or {}).get("source_form")
    try:
        page_int = int(page or 0)
    except Exception:
        page_int = 0
    return _format_source_ref(page_int, source_form)


def _source_form_from_value_item(item: dict[str, Any], metric: dict[str, Any] | None) -> str:
    return _clean_source_form(item.get("source_form") or (metric or {}).get("source_form"))


def _source_from_period_values(values: dict[str, dict[str, Any]], periods: list[str]) -> str:
    refs: dict[int, str] = {}
    for period in periods:
        item = values.get(period) or {}
        _add_source_ref(refs, item.get("page"), item.get("source_form"))
    return ", ".join(_format_source_ref(page, refs.get(page)) for page in sorted(refs))


def _display_unit(unit: object) -> str:
    raw = str(unit or "").replace("\u00a0", " ").strip()
    if not raw:
        return ""
    normalized = re.sub(r"\s+", " ", raw).lower()
    has_usd = "$" in raw or bool(re.search(r"\b(?:usd|us\$|u\.s\. dollars?|dollars?)\b", normalized))
    has_million = bool(
        re.search(r"\b(?:million|millions|mln)\b", normalized)
        or "млн" in normalized
    )
    if has_usd and has_million:
        return "млн. $"
    return raw


def _unit_for(metric: dict[str, Any] | None, meta: dict[str, Any]) -> str:
    if not metric:
        return ""
    for pv in metric.get("period_values") or []:
        if isinstance(pv, dict) and pv.get("unit"):
            return _display_unit(pv.get("unit"))
    return _display_unit(metric.get("unit") or meta.get("unit") or "")


def _is_die_table_unit(unit: str) -> bool:
    normalized = re.sub(r"[\s.]+", "", str(unit or "").lower())
    if not normalized:
        return False
    if normalized in {"%", "пп", "п.п", "п.п.", "раз", "x", "times", "bps", "bp"}:
        return False
    if normalized.endswith("%") or "percentagepoint" in normalized:
        return False
    return True


def _die_table_unit(
    metrics: list[dict[str, Any]],
    meta: dict[str, Any],
    periods: list[str],
    *,
    die_period_mode: str,
    die_period_aliases: dict[str, str] | None,
) -> str:
    counts: dict[str, int] = {}

    def add(unit: object) -> None:
        display = _display_unit(unit)
        if not _is_die_table_unit(display):
            return
        counts[display] = counts.get(display, 0) + 1

    for metric in metrics:
        vals = _period_values(
            metric,
            meta,
            die=True,
            die_period_mode=die_period_mode,
            die_period_aliases=die_period_aliases,
        )
        for period in periods:
            item = vals.get(period) or {}
            if _value_is_present(item.get("value")):
                add(item.get("unit") or metric.get("unit") or meta.get("unit"))

    if counts:
        return sorted(counts.items(), key=lambda item: (-item[1], item[0]))[0][0]

    fallback = _display_unit(meta.get("unit"))
    return fallback if _is_die_table_unit(fallback) else ""


_DIE_SECTION_UNIT_RE = re.compile(
    r"\s*\([^)]*(?:\$|usd|us\$|rub|rmb|cny|yuan|руб|рубл|тыс|млн|млрд|million|millions|billion|billions|bn)[^)]*\)",
    re.IGNORECASE,
)


def _die_section_label(label: object, unit: str) -> str:
    text = str(label or "")
    if not text:
        return ""
    replacement = f" ({unit})" if unit else ""
    return _DIE_SECTION_UNIT_RE.sub(replacement, text, count=1).strip()


_CONSOLIDATED_GROUP_DIMENSIONS = {"consolidated_group", "consolidation_group", "group_structure"}


def _is_consolidated_group_metric(metric: dict[str, Any]) -> bool:
    dimension = str(metric.get("dimension") or "").strip().lower()
    return dimension in _CONSOLIDATED_GROUP_DIMENSIONS


def _is_segment_metric(metric: dict[str, Any]) -> bool:
    segment = str(metric.get("segment") or "").strip()
    if not segment or segment.upper() == "TOTAL":
        return False
    if metric.get("segment_fallback") == "main_metric_total":
        return False
    return not _is_consolidated_group_metric(metric)


def _previous_comparable_period(period: str, available_periods: list[str] | None = None) -> str | None:
    info = _parse_period(period)
    if not info:
        return None
    if info.kind == "quarter":
        return f"{info.family} {info.year - 1}"
    if info.kind == "fy":
        return f"FY {info.year - 1}"
    if info.kind in {"months", "month_range"}:
        return f"{info.family} {info.year - 1}"
    if info.kind == "rolling_months":
        candidates: list[tuple[int, str]] = []
        current_dt = _rolling_period_date(period)
        for candidate in available_periods or []:
            candidate_info = _parse_period(candidate)
            if (
                not candidate_info
                or candidate_info.kind != "rolling_months"
                or candidate_info.family != info.family
                or candidate_info.year != info.year - 1
            ):
                continue
            candidate_dt = _rolling_period_date(candidate)
            if candidate_dt and current_dt:
                distance = abs(
                    (candidate_dt.month * 31 + candidate_dt.day)
                    - (current_dt.month * 31 + current_dt.day)
                )
            else:
                distance = 999
            candidates.append((distance, candidate))
        if candidates:
            return sorted(candidates)[0][1]
        if current_dt:
            try:
                return f"{info.family} {current_dt.replace(year=current_dt.year - 1):%d.%m.%Y}"
            except ValueError:
                return f"{info.family} {current_dt.replace(year=current_dt.year - 1, day=28):%d.%m.%Y}"
    return None


def _change_from_values(vals: dict[str, dict[str, Any]], period: str) -> Any:
    prev_period = _previous_comparable_period(period, list(vals.keys()))
    if not prev_period:
        return ""
    curr = _to_number(vals.get(period, {}).get("value"))
    prev = _to_number(vals.get(prev_period, {}).get("value"))
    if curr is None or prev is None or prev == 0:
        return ""
    return (curr - prev) / prev


def _die_segment_title(enum: str, label: str) -> str:
    if enum in _DIE_SEGMENT_TITLES:
        return _DIE_SEGMENT_TITLES[enum]
    clean = (label or enum).strip()
    return f"Структура {clean[:1].lower() + clean[1:]}" if clean else "Структура показателя"


def _pseudo_metric(enum: str, idx: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    if enum == "SALES_REVENUE_LTM":
        return idx.get("SALES_REVENUE")
    if enum == "EBITDA_LTM":
        return idx.get("EBITDA")
    if enum == "NET_INCOME_LTM":
        return idx.get("NET_INCOME")
    if enum == "NET_DEBT_FOR_RISKS":
        return idx.get("NET_DEBT")
    return idx.get(enum)


def _wide_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    periods = _all_periods(metrics, meta)
    columns = [
        {"key": "label", "label": "Метрика", "kind": "text"},
        {"key": "unit", "label": "Ед. изм.", "kind": "text"},
        *[{"key": p, "label": _display_period(p), "kind": "value"} for p in periods],
    ]
    rows: list[dict[str, Any]] = []
    preserve_empty = bool((plan.config or {}).get("preserve_empty", True))
    for row_cfg in (plan.config or {}).get("rows") or []:
        enum = str(row_cfg.get("enum") or "").strip().upper()
        metric = _pseudo_metric(enum, idx)
        values = _period_values(metric, meta)
        if not preserve_empty and not values:
            continue
        row = {
            "kind": "metric",
            "label": row_cfg.get("label") or (metric or {}).get("name_ru") or enum,
            "enum": enum,
            "unit": _unit_for(metric, meta),
            "source": _source(metric),
            "values": {p: values.get(p, {}).get("value", MISSING_TEXT if preserve_empty else "") for p in periods},
        }
        rows.append(row)
    return _table(plan, columns, rows, periods)


def _baifp_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    periods = _all_periods(metrics, meta)
    bank_name = _company_name_from_meta(meta)
    columns = [
        {"key": "label", "label": "название показателя", "kind": "text"},
        {"key": "form", "label": "форма", "kind": "text"},
        {"key": "bank", "label": "название банка", "kind": "text"},
        {"key": "period", "label": "период (год)", "kind": "text"},
        {"key": "value", "label": "Значение", "kind": "value"},
        {"key": "unit", "label": "Единицы измерения", "kind": "text"},
    ]
    rows: list[dict[str, Any]] = []
    for row_cfg in (plan.config or {}).get("rows") or []:
        enum = str(row_cfg.get("enum") or "").strip().upper()
        metric = _pseudo_metric(enum, idx)
        vals = _period_values(metric, meta)
        for period in periods:
            value = vals.get(period, {}).get("value")
            rows.append(
                {
                    "kind": "metric",
                    "enum": enum,
                    "label": row_cfg.get("label") or (metric or {}).get("name_ru") or enum,
                    "form": _source_form_from_value_item(vals.get(period) or {}, metric),
                    "bank": bank_name,
                    "period": _legacy_period_label(period),
                    "value": value if _value_is_present(value) else MISSING_TEXT,
                    "unit": _unit_for(metric, meta),
                    "values": {},
                }
            )
    return _table(plan, columns, rows, periods)


def _cib_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    rows: list[dict[str, Any]] = []
    for row_cfg in (plan.config or {}).get("rows") or []:
        enum = str(row_cfg.get("enum") or "").strip().upper()
        metric = _pseudo_metric(enum, idx)
        vals = _period_values(metric, meta)
        if not vals:
            continue
        unit, currency = _unit_currency(_unit_for(metric, meta))
        label = row_cfg.get("label") or (metric or {}).get("name_ru") or enum
        for period in _paired_period_order(_compact_periods(list(vals))):
            item = vals[period]
            value = item.get("value")
            rows.append(
                {
                    "kind": "metric",
                    "enum": enum,
                    "label": label,
                    "source_note": _source_from_value_item(item, metric),
                    "period": _legacy_period_label(period),
                    "value": _to_number(value) if _to_number(value) is not None else value,
                    "unit": unit if _value_is_present(value) else "",
                    "currency": currency if _value_is_present(value) else "",
                    "values": {},
                }
            )
    rows.sort(key=lambda r: str(r.get("label") or ""))
    for i, row in enumerate(rows, start=1):
        row["n"] = i
    columns = [
        {"key": "n", "label": "№", "kind": "number"},
        {"key": "label", "label": "Название на русском", "kind": "text"},
        {"key": "source_note", "label": "Источник (отчет, примечание)", "kind": "text"},
        {"key": "period", "label": "Период отчетности", "kind": "text"},
        {"key": "value", "label": "Значение", "kind": "value"},
        {"key": "unit", "label": "Единицы измерения", "kind": "text"},
        {"key": "currency", "label": "Валюта", "kind": "text"},
    ]
    return _table(plan, columns, rows, [])


def _risk_period_pairs(periods: list[str]) -> list[tuple[str, str]]:
    grouped: dict[tuple[str, int, int], dict[int, str]] = {}
    for period in periods:
        info = _parse_period(period)
        if not info:
            continue
        if info.kind == "fy":
            family = ("FY", 0, 0)
        elif info.kind == "quarter":
            quarter = int(info.family.replace("Q", ""))
            family = (info.family, 1, quarter)
        elif info.kind == "months":
            month = int(info.family.replace("M", ""))
            family = (info.family, 2, month)
        elif info.kind == "date":
            family = (f"DATE:{period[:5]}", 4, int(period[3:5]) * 100 + int(period[:2]))
        else:
            continue
        grouped.setdefault(family, {})[info.year] = period
    pairs: list[tuple[tuple[str, int, int], str, str]] = []
    for family, by_year in grouped.items():
        years = sorted(by_year)
        for prev, curr in zip(years, years[1:], strict=False):
            if curr == prev + 1:
                pairs.append((family, by_year[prev], by_year[curr]))
    pairs.sort(key=lambda item: (item[0][1], item[0][2], _period_sort_key(item[2])))
    return [(prev, curr) for _, prev, curr in pairs]


def _risk_unpaired_periods(periods: list[str], pairs: list[tuple[str, str]]) -> list[str]:
    paired = {period for pair in pairs for period in pair}
    return [period for period in sorted(periods, key=_period_sort_key) if period not in paired]


def _paired_period_order(periods: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for prev, curr in _risk_period_pairs(periods):
        for period in (prev, curr):
            if period not in seen:
                seen.add(period)
                out.append(period)
    for period in sorted(periods, key=_period_sort_key):
        if period not in seen:
            out.append(period)
    return out


def _delta(curr: Any, prev: Any, *, pct: bool) -> Any:
    c = _to_number(curr)
    p = _to_number(prev)
    if c is None or p is None or (pct and c == 0):
        return ""
    if pct:
        return (c - p) / c
    return c - p


def _risks_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    periods = _all_periods(metrics, meta)
    pairs = _risk_period_pairs(periods)
    unpaired_periods = _risk_unpaired_periods(periods, pairs)
    columns = [
        {"key": "label_en", "label": "финансовые коэффициенты", "kind": "text"},
        {"key": "blank1", "label": "", "kind": "text"},
    ]
    for pair_idx, (prev, curr) in enumerate(pairs):
        columns.extend(
            [
                {"key": prev, "label": _legacy_period_label(prev), "kind": "value"},
                {"key": curr, "label": _legacy_period_label(curr), "kind": "value"},
                {"key": f"delta:{prev}:{curr}", "label": "динамика", "kind": "value"},
                {"key": f"delta_pct:{prev}:{curr}", "label": "динамика", "kind": "percent"},
            ]
        )
        if pair_idx != len(pairs) - 1:
            columns.append({"key": f"spacer:{pair_idx}", "label": "", "kind": "text"})
    if not pairs:
        columns.extend({"key": p, "label": _legacy_period_label(p), "kind": "value"} for p in periods)
    else:
        columns.extend({"key": p, "label": _legacy_period_label(p), "kind": "value"} for p in unpaired_periods)
    columns.extend(
        [
            {"key": "blank2", "label": "", "kind": "text"},
        ]
    )

    rows: list[dict[str, Any]] = []
    for row_cfg in (plan.config or {}).get("rows") or []:
        enum = str(row_cfg.get("enum") or "").strip().upper()
        metric = _pseudo_metric(enum, idx)
        vals = _period_values(metric, meta)
        row: dict[str, Any] = {
            "kind": "metric",
            "enum": enum,
            "label_en": row_cfg.get("label_en") or enum,
            "blank1": "",
            "blank2": "",
            "values": {},
            "value_kind": row_cfg.get("value_kind") or "number",
        }
        if pairs:
            for prev, curr in pairs:
                prev_value = vals.get(prev, {}).get("value", MISSING_TEXT)
                curr_value = vals.get(curr, {}).get("value", MISSING_TEXT)
                row["values"][prev] = prev_value
                row["values"][curr] = curr_value
                row["values"][f"delta:{prev}:{curr}"] = _delta(curr_value, prev_value, pct=False)
                row["values"][f"delta_pct:{prev}:{curr}"] = "" if row_cfg.get("no_delta_pct") else _delta(curr_value, prev_value, pct=True)
            for period in unpaired_periods:
                row["values"][period] = vals.get(period, {}).get("value", MISSING_TEXT)
        else:
            for period in periods:
                row["values"][period] = vals.get(period, {}).get("value", MISSING_TEXT)
        rows.append(row)
    return _table(plan, columns, rows, periods)


def _die_fraction_for_percent_row(value: Any, unit: object) -> Any:
    """Значение строки `is_percent` — ДОЛЯ: рендереры умножат её на 100 сами.

    Оба рендерера ДИЭ читают долю: `web/src/lib/formatted-table.ts:76`
    (`value * 100`) и `_write_die_sheet` ниже (`number_format = "0%"`). Пока
    реестр держал доходности долями, строки `kind == "metric"` можно было
    отдавать как есть; после перевода семейства процентов в проценты то же
    место показывало ROE 24,6 как «2460 %».

    Различитель — ЕДИНИЦА строки, а не список enum'ов: список стал бы ещё
    одним местом, где живёт знание о единице показателя, и разъехался бы с
    реестром. Единицу расчётной строки на всех путях назначает
    `unit_for_calculated`, поэтому «%» здесь означает «число уже в процентах».
    Всё прочее (доля, деньги) отдаём как есть — как это место вело себя до
    реестровой правки.
    """
    if str(unit or "").strip() != "%":
        return value
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return value
    return value / 100.0


def _die_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    die_period_mode = _die_period_mode(metrics, meta)
    die_period_aliases = _die_period_aliases(metrics, meta, mode=die_period_mode)
    periods = _all_periods(
        metrics,
        meta,
        die=True,
        die_period_mode=die_period_mode,
        die_period_aliases=die_period_aliases,
    )
    periods = [period for period in periods if _is_die_output_period(period)]
    unit_label = _die_table_unit(
        metrics,
        meta,
        periods,
        die_period_mode=die_period_mode,
        die_period_aliases=die_period_aliases,
    )
    columns = [
        {"key": "label", "label": unit_label, "kind": "text"},
        *[{"key": p, "label": _display_period(p, style="die"), "kind": "value"} for p in periods],
        {"key": "source", "label": "Источник", "kind": "text"},
    ]

    def calc_change(enum: str, period: str, *, ytd: bool) -> Any:
        info = _parse_period(period)
        if not info:
            return ""
        vals = _period_values(
            _pseudo_metric(enum, idx),
            meta,
            die=True,
            die_period_mode=die_period_mode,
            die_period_aliases=die_period_aliases,
        )
        if info.kind == "quarter":
            cmp_period = f"Q4 {info.year - 1}" if ytd else f"{info.family} {info.year - 1}"
        elif info.kind == "fy":
            cmp_period = f"FY {info.year - 1}"
        elif ytd:
            return ""
        else:
            cmp_period = _previous_comparable_period(period, list(vals.keys())) or ""
        if die_period_aliases:
            cmp_period = die_period_aliases.get(cmp_period, cmp_period)
        curr = vals.get(period, {}).get("value")
        prev = vals.get(cmp_period, {}).get("value")
        c = _to_number(curr)
        p = _to_number(prev)
        if c is None or p is None or p == 0:
            return ""
        return (c - p) / p

    rows: list[dict[str, Any]] = []
    for row_cfg in (plan.config or {}).get("rows") or []:
        kind = str(row_cfg.get("kind") or "")
        enum = str(row_cfg.get("enum") or "").strip().upper()
        base_enum = str(row_cfg.get("base_enum") or "").strip().upper()
        row = {
            "kind": "section" if kind == "section" else "spacer" if kind == "spacer" else "metric",
            "die_kind": kind,
            "enum": enum,
            "label": _die_section_label(row_cfg.get("title"), unit_label) if kind == "section" else row_cfg.get("title") or "",
            "is_percent": bool(row_cfg.get("is_percent")),
            "values": {},
        }
        if kind in {"section", "spacer"}:
            rows.append(row)
            continue
        metric_for_row = _pseudo_metric(enum, idx)
        vals_for_row = _period_values(
            metric_for_row,
            meta,
            die=True,
            die_period_mode=die_period_mode,
            die_period_aliases=die_period_aliases,
        )
        if kind == "metric":
            row["source"] = _source_from_period_values(vals_for_row, periods)
        for period in periods:
            if kind == "metric":
                _item = vals_for_row.get(period, {})
                value = _item.get("value")
                if not _value_is_present(value):
                    row["values"][period] = ""
                elif row["is_percent"]:
                    row["values"][period] = _die_fraction_for_percent_row(value, _item.get("unit"))
                else:
                    row["values"][period] = value
            elif kind == "yoy":
                row["values"][period] = calc_change(base_enum, period, ytd=False)
            elif kind == "ytd":
                row["values"][period] = calc_change(base_enum, period, ytd=True)
            elif kind in {"margin", "percent_of_revenue"}:
                # Margins/ratios are stored in percent (e.g. EBIT_MARGIN = 65.6).
                # The DIE renderers (web fmtFormattedCell, xlsx "0%") multiply
                # is_percent cells by 100, so scale to a fraction here to avoid
                # double-scaling (65.6 -> 0.656 -> 65.6%).
                _margin_value = vals_for_row.get(period, {}).get("value")
                row["values"][period] = (
                    _margin_value / 100.0
                    if isinstance(_margin_value, (int, float))
                    and not isinstance(_margin_value, bool)
                    else (_margin_value or "")
                )
        rows.append(row)

    segment_metrics_by_enum: dict[str, list[dict[str, Any]]] = {}
    for metric in metrics:
        enum = str(metric.get("enum") or "").strip().upper()
        if not enum or not _is_segment_metric(metric):
            continue
        segment_metrics_by_enum.setdefault(enum, []).append(metric)

    appended_any_segments = False
    for row_cfg in (plan.config or {}).get("rows") or []:
        if str(row_cfg.get("kind") or "") != "metric":
            continue
        enum = str(row_cfg.get("enum") or "").strip().upper()
        segment_metrics = segment_metrics_by_enum.get(enum) or []
        if not segment_metrics:
            continue

        base_label = row_cfg.get("title") or row_cfg.get("label") or row_cfg.get("label_en") or (idx.get(enum) or {}).get("name_ru") or enum
        if not appended_any_segments:
            rows.append({"kind": "spacer", "die_kind": "spacer", "label": "", "values": {}})
            appended_any_segments = True
        rows.append(
            {
                "kind": "section",
                "die_kind": "segment_section",
                "enum": enum,
                "label": _die_segment_title(enum, str(base_label)),
                "is_percent": False,
                "values": {},
            }
        )

        seen_segments: set[tuple[str, str]] = set()
        for metric in segment_metrics:
            segment = str(metric.get("segment") or "").strip()
            dimension = str(metric.get("dimension") or "").strip()
            key = (segment.lower(), dimension.lower())
            if key in seen_segments:
                continue
            seen_segments.add(key)

            vals = _period_values(
                metric,
                meta,
                die=True,
                die_period_mode=die_period_mode,
                die_period_aliases=die_period_aliases,
            )
            segment_values: dict[str, Any] = {}
            yoy_values: dict[str, Any] = {}
            for period in periods:
                value = vals.get(period, {}).get("value")
                segment_values[period] = value if _value_is_present(value) else ""
                yoy_values[period] = _change_from_values(vals, period)

            rows.append(
                {
                    "kind": "metric",
                    "die_kind": "segment",
                    "enum": enum,
                    "segment": segment,
                    "dimension": dimension,
                    "label": segment,
                    "source": _source_from_period_values(vals, periods),
                    "bold": True,
                    "is_percent": False,
                    "values": segment_values,
                }
            )
            rows.append(
                {
                    "kind": "metric",
                    "die_kind": "segment_yoy",
                    "enum": enum,
                    "segment": segment,
                    "dimension": dimension,
                    "label": "...YoY",
                    "is_percent": True,
                    "values": yoy_values,
                }
            )
    return _table(plan, columns, rows, periods)


def _dgr_table(plan: RolePlan, metrics: list[dict[str, Any]], meta: dict[str, Any]) -> dict[str, Any]:
    idx = _metric_index(metrics)
    periods = _all_periods(metrics, meta)
    columns = [
        {"key": "enum", "label": "MetricEnum", "kind": "text", "hidden": True},
        {"key": "label", "label": "Название на русском", "kind": "text"},
        *[{"key": p, "label": _legacy_period_label(p), "kind": "value"} for p in periods],
    ]
    rows: list[dict[str, Any]] = []
    for row_cfg in (plan.config or {}).get("rows") or []:
        enum = str(row_cfg.get("enum") or "").strip().upper()
        metric = _pseudo_metric(enum, idx)
        vals = _period_values(metric, meta)
        values: dict[str, Any] = {}
        for period in periods:
            if not enum or not row_cfg.get("is_available"):
                values[period] = ""
            else:
                values[period] = vals.get(period, {}).get("value", MISSING_TEXT)
        rows.append(
            {
                "kind": "spacer" if not row_cfg.get("label") and not enum else "metric",
                "enum": enum,
                "label": row_cfg.get("label") or "",
                "indent": int(row_cfg.get("indent") or 0),
                "bold": bool(row_cfg.get("bold")),
                "highlight": bool(row_cfg.get("is_highlighted") or row_cfg.get("is_available")),
                "available": bool(row_cfg.get("is_available")),
                "values": values,
            }
        )
    return _table(plan, columns, rows, periods)


def _table(plan: RolePlan, columns: list[dict[str, Any]], rows: list[dict[str, Any]], periods: list[str]) -> dict[str, Any]:
    return {
        "role": plan.role,
        "role_label": plan.role_label,
        "company_type": plan.company_type,
        "format": plan.format,
        "sheet_name": "MGlob" if plan.format in {"dgr", "mglob"} else "Данные",
        "columns": columns,
        "rows": rows,
        "periods": periods,
    }


def build_formatted_table(
    *,
    plan: RolePlan,
    metrics: list[dict[str, Any]],
    meta: dict[str, Any] | None = None,
) -> dict[str, Any] | None:
    if plan.role == ALL_EXTRACTION_ROLE or not plan.config:
        return None
    meta = _with_period_context(metrics, meta or {})
    fmt = plan.format
    if fmt == "wide":
        return _wide_table(plan, metrics, meta)
    if fmt == "risks":
        return _risks_table(plan, metrics, meta)
    if fmt == "baifp":
        return _baifp_table(plan, metrics, meta)
    if fmt == "cib":
        return _cib_table(plan, metrics, meta)
    if fmt in {"die_nf", "die_bank"}:
        return _die_table(plan, metrics, meta)
    if fmt in {"dgr", "mglob"}:
        return _dgr_table(plan, metrics, meta)
    return None


def _row_cell(row: dict[str, Any], column: dict[str, Any]) -> Any:
    key = column["key"]
    if key in row:
        return row.get(key)
    return (row.get("values") or {}).get(key, "")


def _xlsx_buf(wb: Any) -> BytesIO:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _xlsx_number(value: Any) -> int | float | Any:
    number = _to_number(value)
    if number is None:
        return value
    return int(number) if number.is_integer() else number


def _set_widths(ws: Any, widths: list[float]) -> None:
    from openpyxl.utils import get_column_letter

    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _autosize_plain_columns(ws: Any, *, max_width: int = 60, padding: int = 2) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx in range(1, ws.max_column + 1):
        longest = 0
        for row_idx in range(1, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is None:
                continue
            longest = max(longest, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_width, max(9, longest + padding))


def _write_plain_xlsx(table: dict[str, Any]) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(table.get("sheet_name") or "Данные")[:31]

    columns = list(table.get("columns") or [])
    rows = list(table.get("rows") or [])
    fmt = str(table.get("format") or "")

    ws.append([c.get("label") or "" for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([_xlsx_number(_row_cell(row, column)) for column in columns])

    ws.freeze_panes = "A2"
    if fmt == "baifp":
        _set_widths(ws, [48, 10, 16, 14, 10, 19])
        value_col = 5
    elif fmt == "cib":
        _set_widths(ws, [5, 60, 30, 19, 10, 19, 8])
        value_col = 5
    else:
        _autosize_plain_columns(ws)
        value_col = 0

    if value_col:
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=value_col)
            if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
                continue
            cell.number_format = "#,##0" if isinstance(cell.value, int) else "#,##0.00"

    return _xlsx_buf(wb)


def _write_risks_xlsx(table: dict[str, Any]) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(table.get("sheet_name") or "Данные")[:31]

    columns = list(table.get("columns") or [])
    rows = list(table.get("rows") or [])

    ws.append([c.get("label") or "" for c in columns])
    ws.append([""] * len(columns))
    for row in rows:
        values = []
        for column in columns:
            value = _row_cell(row, column)
            if value == MISSING_TEXT:
                value = None
            values.append(_xlsx_number(value))
        ws.append(values)

    ws.freeze_panes = "C3"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 8

    fill_green = PatternFill("solid", fgColor="92D050")
    fill_yellow = PatternFill("solid", fgColor="FFF2CC")
    fill_gray = PatternFill("solid", fgColor="D9D9D9")
    fill_dyn = PatternFill("solid", fgColor="E7E6E6")
    font_header = Font(bold=True)
    font_green = Font(bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.alignment = align_center
        if col_idx == 1:
            cell.fill = fill_green
            cell.font = font_green
            cell.alignment = align_left
        elif col_idx >= 3 and cell.value not in (None, ""):
            cell.fill = fill_yellow
        ws.cell(row=2, column=col_idx).font = font_header
        ws.cell(row=2, column=col_idx).alignment = align_center

    data_start = 3
    last_row = ws.max_row
    blank_col = ws.max_column - 1 if ws.max_column > 1 else ws.max_column
    for row_idx, row in enumerate(rows, start=data_start):
        value_kind = str(row.get("value_kind") or "number")
        fmt_year = "0.00%" if value_kind == "percent" else '0.00;-0.00;"-"' if value_kind == "ratio" else '#,##0;-#,##0;"-"'
        for col_idx, column in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(
                horizontal="left" if col_idx == 1 else "center" if col_idx in {2, blank_col} else "right",
                vertical="center",
                wrap_text=col_idx == 1,
            )
            if col_idx == 1:
                cell.fill = fill_gray
                cell.font = Font(bold=True)
            if str(ws.cell(row=1, column=col_idx).value or "") == "динамика":
                cell.fill = fill_dyn
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                if column.get("kind") == "percent":
                    cell.number_format = "0.00%"
                else:
                    cell.number_format = fmt_year

    for row_idx in range(1, last_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            left = medium if col_idx in {1, 3} else thin
            right = medium if col_idx == ws.max_column else thin
            top = medium if row_idx == 1 else thin
            bottom = medium if row_idx in {2, last_row} else thin
            header = str(ws.cell(row=1, column=col_idx).value or "")
            prev_header = str(ws.cell(row=1, column=col_idx - 1).value or "") if col_idx > 1 else ""
            if header == "динамика" and prev_header != header:
                left = medium
            ws.cell(row=row_idx, column=col_idx).border = Border(
                left=left,
                right=right,
                top=top,
                bottom=bottom,
            )

    for col_idx, column in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        key = str(column.get("key") or "")
        label = str(column.get("label") or "")
        if col_idx == 1:
            width = 30
        elif col_idx == 2:
            width = 1
        elif key.startswith("delta_pct:"):
            width = 10
        elif key.startswith("delta:"):
            width = 14
        elif key.startswith("spacer:") or key == "blank2":
            width = 2
        elif key == "label_ru":
            width = 90
        elif label:
            width = 12
        else:
            width = 2
        ws.column_dimensions[letter].width = width

    return _xlsx_buf(wb)


def _write_die_xlsx(table: dict[str, Any]) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(table.get("sheet_name") or "Данные")[:31]

    columns = list(table.get("columns") or [])
    rows = list(table.get("rows") or [])
    fmt = str(table.get("format") or "")
    is_bank = fmt == "die_bank"

    ws.append([c.get("label") or "" for c in columns])
    for row in rows:
        ws.append([_xlsx_number(_row_cell(row, column)) for column in columns])

    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 22
    header_fill = PatternFill("solid", fgColor="D9D9D9" if is_bank else "D9EAD3")
    metric_fill = PatternFill("solid", fgColor="F2F2F2" if is_bank else "D9D9D9")
    calc_fill = PatternFill("solid", fgColor="E2EFDA" if is_bank else "EFEFEF")
    white_fill = PatternFill("solid", fgColor="FFFFFF")
    section_side = Side(style="medium", color="000000")
    section_border = Border(top=section_side, bottom=section_side)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, start=2):
        die_kind = str(row.get("die_kind") or "")
        if row.get("kind") == "section":
            fill = white_fill
        elif row.get("kind") == "spacer":
            fill = white_fill
        elif die_kind in {"yoy", "ytd", "margin", "percent_of_revenue", "segment_yoy"}:
            fill = calc_fill
        else:
            fill = metric_fill

        if row.get("kind") == "spacer":
            ws.row_dimensions[row_idx].height = 8
        elif die_kind == "segment_section":
            ws.row_dimensions[row_idx].height = 22

        for col_idx, column in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            key = str(column.get("key") or "")
            cell.fill = fill
            if die_kind == "segment_section":
                cell.border = section_border
            cell.alignment = Alignment(
                horizontal="left" if col_idx == 1 or key == "source" else "right",
                vertical="center",
                wrap_text=col_idx == 1 or key == "source",
            )
            if row.get("kind") == "section" or row.get("bold") or (is_bank and col_idx == ws.max_column):
                cell.font = Font(bold=True)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                if column.get("kind") == "percent" or row.get("is_percent") or die_kind in {"yoy", "ytd", "margin", "percent_of_revenue", "segment_yoy"}:
                    cell.number_format = "0%"
                else:
                    cell.number_format = "#,##0.0"

    widths = []
    for column in columns:
        key = str(column.get("key") or "")
        if key == "label":
            widths.append(36 if is_bank else 28)
        elif key == "source":
            widths.append(16)
        else:
            widths.append(11 if is_bank else 12)
    _set_widths(ws, widths)
    return _xlsx_buf(wb)


def _write_dgr_xlsx(table: dict[str, Any]) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(table.get("sheet_name") or "MGlob")[:31]

    columns = list(table.get("columns") or [])
    rows = list(table.get("rows") or [])
    ws.append([c.get("label") or "" for c in columns])
    for row in rows:
        values = []
        for column in columns:
            values.append(_xlsx_number(_row_cell(row, column)))
        ws.append(values)

    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].hidden = True
    ws.row_dimensions[1].height = 24
    _set_widths(ws, [13, 70] + [14] * max(0, ws.max_column - 2))

    header_fill = PatternFill("solid", fgColor="D9EAD3")
    available_fill = PatternFill("solid", fgColor="A9D18E")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_idx, row in enumerate(rows, start=2):
        if row.get("kind") == "spacer":
            ws.row_dimensions[row_idx].height = 8
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
                indent=int(row.get("indent") or 0) if col_idx == 2 else 0,
            )
            if row.get("bold"):
                cell.font = Font(bold=True)
            if row.get("highlight") or row.get("available"):
                cell.fill = available_fill
            if col_idx >= 3 and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0" if float(cell.value).is_integer() else "#,##0.00"

    return _xlsx_buf(wb)


def write_generic_wide_xlsx(metrics: list[dict[str, Any]], meta: dict[str, Any] | None = None) -> BytesIO:
    import openpyxl
    from openpyxl.styles import Font

    meta = _with_period_context(metrics, meta or {})
    periods = _all_periods(metrics, meta)
    columns = [
        {"key": "label", "label": "Метрика", "kind": "text"},
        {"key": "unit", "label": "Ед. изм.", "kind": "text"},
        *[{"key": p, "label": _legacy_period_label(p), "kind": "value"} for p in periods],
        {"key": "source", "label": "Источник", "kind": "text"},
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append([c["label"] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    names = {str(m.get("enum") or "").strip(): str(m.get("name_ru") or "").strip()
             for m in metrics if m.get("name_ru")}
    for metric in metrics:
        vals = _period_values(metric, meta)
        label = metric.get("name_ru") or metric.get("enum") or ""
        segment = str(metric.get("segment") or "").strip()
        if segment and segment.upper() != "TOTAL" and not _is_consolidated_group_metric(metric):
            label = f"{label} / {segment}"
        source = str(metric.get("source_text") or "").strip() or _source(metric)
        if not source and metric.get("is_calculated"):
            source = _calculated_source(metric, names)
        row = [label, _unit_for(metric, meta)]
        for period in periods:
            row.append(_xlsx_number(vals.get(period, {}).get("value", "")))
        row.append(source)
        ws.append(row)

    ws.freeze_panes = "A2"
    _set_widths(ws, [60, 11] + [9] * len(periods) + [18])
    for col_idx in range(3, ws.max_column + 1):
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0.00"
    return _xlsx_buf(wb)


def write_formatted_xlsx(table: dict[str, Any]) -> BytesIO:
    fmt = str(table.get("format") or "")
    if fmt == "risks":
        return _write_risks_xlsx(table)
    if fmt in {"die_nf", "die_bank"}:
        return _write_die_xlsx(table)
    if fmt in {"dgr", "mglob"}:
        return _write_dgr_xlsx(table)
    return _write_plain_xlsx(table)

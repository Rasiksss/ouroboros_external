"""Сборка Excel по извлечённым метрикам (все периоды + сегменты + группа).

Один источник правды для двух потребителей:
  - инлайн-тул `fe_extract_metrics` (llm/tools.py) — как и раньше;
  - bg-mailbox фонового извлечения воркером (#84, orchestrator/agent.py) —
    metric-события уже текут через его очередь, бэк собирает тот же файл и
    кладёт в «Источники», чтобы подхват честно ссылался на Excel.

Чистая функция: события → bytes xlsx. Naming/сохранение — на вызывающем.
"""
from __future__ import annotations

import io
import logging

log = logging.getLogger(__name__)


def build_metrics_xlsx(
    metrics: list[dict],
    segment_rows: list[dict] | None = None,
    group_rows: list[dict] | None = None,
    passport: list[tuple[str, str]] | None = None,
    company_type: str = "",
) -> bytes:
    """Собирает книгу: лист «Показатели» (все периоды), опц. «Сегменты», «Группа»,
    опц. «Паспорт прогона».
    Формат листов 1:1 с историческим инлайн-блоком fe_extract_metrics.

    `passport` (#155) — пары «подпись → значение» с отпечатком реестра, по
    которому посчитан прогон; готовит их `skills.versioning.passport_rows`.
    Лист нужен потому, что расхождение вчерашней выгрузки с сегодняшней человек
    видит В ФАЙЛЕ, а не в журнале на сервере: без признака версии правил объяснить
    «почему поехало» ему нечем. Пусто — листа нет: пустой паспорт хуже
    отсутствующего, он читается как «правила не менялись».
    """
    import openpyxl

    segment_rows = segment_rows or []
    group_rows = group_rows or []
    plabels = sorted({pv.get("period") for m in metrics
                      for pv in (m.get("period_values") or []) if pv.get("period")})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Показатели"
    # «Стр. PDF» — физическая страница файла, на которой движок НАШЁЛ число
    # (#157). Тот же номер, что стоит ссылкой в отчёте анализа и по которому
    # прыгает панель документа: выгрузка обязана давать возможность проверить
    # число в исходнике, не листая отчёт руками. Колонка последняя — порядок
    # уже существующих столбцов трогать нельзя, на него смотрят глазами каждый
    # день. Страницы нет (значение посчитано формулой) — пусто, а не ноль:
    # ноль читался бы как номер страницы.
    ws.append(["Показатель", "enum"] + plabels + ["Ед.", "Стр. PDF"])
    for m in metrics:
        pvmap = {pv.get("period"): pv.get("value") for pv in (m.get("period_values") or [])}
        page = m.get("page")
        ws.append([m.get("name_ru") or "", m.get("enum")] + [pvmap.get(p) for p in plabels]
                  + [m.get("unit") or "", page if isinstance(page, int) and page > 0 else ""])
    # Segment breakdown (по сегментам/географии/продуктам) → отдельный лист.
    if segment_rows:
        seg_pl = sorted({pv.get("period") for m in segment_rows
                         for pv in (m.get("period_values") or []) if pv.get("period")})
        ws2 = wb.create_sheet("Сегменты")
        ws2.append(["Показатель", "enum", "Сегмент", "Измерение"] + seg_pl + ["Ед."])
        for m in segment_rows:
            pvmap = {pv.get("period"): pv.get("value") for pv in (m.get("period_values") or [])}
            ws2.append([m.get("name_ru") or "", m.get("enum"), m.get("segment") or "",
                        m.get("dimension") or ""] + [pvmap.get(p) for p in seg_pl] + [m.get("unit") or ""])
    # Состав консолидированной группы → отдельный лист.
    if group_rows:
        ws3 = wb.create_sheet("Группа")
        ws3.append(["Компания", "Доля", "Ед.", "Отношение", "Регистрация", "Деятельность"])
        for g in group_rows:
            ws3.append([g.get("name_ru") or g.get("segment") or "", g.get("value_current"),
                        g.get("unit") or "", g.get("group_relation") or "",
                        g.get("group_registration") or "", g.get("group_activity") or ""])
    # Данные графиков карточки анализа (#157). Столбик графика человек должен
    # уметь проверить руками, а на картинке чисел нет — поэтому спецификация
    # выкладывается числами в книгу. Строится ЗДЕСЬ, из того же списка metrics,
    # по которому собран лист «Показатели», и той же функцией, что кормит
    # карточку: два независимых сборщика — это два расходящихся ответа на один
    # вопрос (ровно дефект эталона: 41 150 в HTML против 32 393 в PDF).
    # Не собрался (нет входов, разошлись единицы) — листа нет, причина в логе.
    #
    # ⚠️ ЗАВЁРНУТО — как и второй вызов этой же функции в `llm/tools.py`, где
    # написано «график не смеет ронять готовый отчёт». Здесь цена выше: падение
    # уносит не приложение к отчёту, а ВСЮ книгу извлечения — единственный
    # выход прогона, который пользователь уносит к себе. Показатель с
    # нечисловым значением («Немодифицированное» в аудиторском мнении — штатный
    # случай реестра) роняет `float()` внутри редьюсера; проверено исполнением:
    # до этой обёртки книга не собиралась вовсе.
    from finengine.skills.analysis.charts import build_charts

    try:
        charts = build_charts(metrics, company_type)
    except Exception:  # noqa: BLE001 — лист графиков не смеет ронять книгу
        log.exception("metrics_xlsx: спецификации графиков не собраны")
        charts = []
    if charts:
        ws5 = wb.create_sheet("Данные графиков")
        ws5.append(["График", "Ед.", "Столбик", "Значение", "Тип", "Источник"])
        for spec in charts:
            for p in spec.get("points") or []:
                ws5.append([spec.get("title") or "", spec.get("unit") or "",
                            p.get("label") or "", p.get("value"), p.get("kind") or "",
                            spec.get("source") or ""])
    # Паспорт прогона — последним листом: он про прогон, а не про данные, и
    # первым открывающимся должен оставаться лист с показателями.
    if passport:
        ws4 = wb.create_sheet("Паспорт прогона")
        ws4.append(["Признак", "Значение"])
        for label, value in passport:
            ws4.append([label, value])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

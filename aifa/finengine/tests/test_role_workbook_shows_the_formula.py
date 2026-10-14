"""В ролевой книге у расчётной строки виден источник — формула с числами.

ЦЕНА (Б5, 19.08.2026). `grep -c formula src/finengine/roles/formatter.py` давал
НОЛЬ на 1787 строк: книга прогона, которую скачивает аналитик, не несла ни
формулы, ни входов. Колонка «Источник» у расчётной строки при этом пустовала —
`_source` собирает ссылки по номеру страницы, а у расчёта страницы нет.

ГРАНИЦА ПРАВКИ НАЗВАНА ВСЛУХ. Починена ОДНА раскладка из пяти — широкая
`write_generic_wide_xlsx`. Специализированные (`die`, `dgr`, `risks`, `plain`)
идут через `write_formatted_xlsx` и строят свои колонки; лезть в них значило бы
переписать форматтер, а не дописать источник. Это осознанный остаток, а не
недосмотр.

ПРОВЕРКА СТОРОЖЕЙ ОТКАТОМ (проделана 19.08):
  1. вернуть `source = ... or _source(metric)` без ветки расчёта → падает
     `test_calculated_row_shows…`;
  2. убрать пометку нуля из `_calculated_source` → падает `test_a_zero…`;
  3. убрать словарь `names` (оставить коды реестра) → падает
     `test_inputs_are_named…`.
"""
from __future__ import annotations

import openpyxl

from finengine.roles.formatter import write_generic_wide_xlsx

PERIOD = "2025"


def _extracted(enum: str, name_ru: str, value: float) -> dict:
    return {
        "enum": enum, "name_ru": name_ru, "unit": "млн ₽", "found": True,
        "is_calculated": False, "formula": "", "page": 7,
        "period_values": [{"period": PERIOD, "value": value, "unit": "млн ₽", "page": 7}],
    }


def _calculated(inputs: list[dict]) -> dict:
    return {
        "enum": "EBITDA", "name_ru": "EBITDA", "unit": "млн ₽", "found": True,
        "is_calculated": True, "page": None,
        "formula": "OPERATING_PROFIT + DEPRECIATION_AMORTIZATION",
        "inputs": inputs,
        "period_values": [{"period": PERIOD, "value": 37850.0, "unit": "млн ₽", "page": None}],
    }


def _source_column(metrics: list[dict]):
    ws = openpyxl.load_workbook(write_generic_wide_xlsx(metrics, {"period_current": PERIOD})).active
    header = [c.value for c in ws[1]]
    col = header.index("Источник") + 1
    return {row[0].value: ws.cell(row=row[0].row, column=col)
            for row in ws.iter_rows(min_row=2, max_col=1)}


def _source_cells(metrics: list[dict]) -> dict[str, str]:
    return {name: (cell.value or "") for name, cell in _source_column(metrics).items()}


ROWS = [
    _extracted("OPERATING_PROFIT", "Операционная прибыль", 28355.0),
    _extracted("DEPRECIATION_AMORTIZATION", "Амортизация", 9495.0),
    _calculated([
        {"enum": "OPERATING_PROFIT", "value": 28355.0, "substituted_zero": False,
         "period": PERIOD},
        {"enum": "DEPRECIATION_AMORTIZATION", "value": 9495.0, "substituted_zero": False,
         "period": PERIOD},
    ]),
]


def test_calculated_row_shows_its_formula_where_the_page_would_be():
    cells = _source_cells(ROWS)
    assert cells["Операционная прибыль"] == "стр. 7", "строку извлечения задели зря"
    assert "Операционная прибыль + Амортизация" in cells["EBITDA"]


def test_inputs_are_named_in_russian_and_carry_their_numbers():
    """Числа — как на экране: разряды неразрывным пробелом, без хвоста `.0`.

    До 20.08 книга печатала `28355.0` и коды реестра, а экран — `28 355` и
    русские имена: два ответа на один вопрос «как это посчитано». Точный вид
    подписи закреплён общей разметкой (`test_calculation_signature_is_one_rule`),
    здесь проверяется, что до КЛЕТКИ книги он доезжает целиком.
    """
    cell = _source_cells(ROWS)["EBITDA"]
    assert "входы за 2025: Операционная прибыль = 28 355" in cell
    assert "Амортизация = 9 495" in cell


def test_the_source_cell_is_text_and_not_a_sheet_formula():
    """Клетка, начинающаяся со знака равенства, — для листа ФОРМУЛА.

    Найдено 20.08 при сведении двух отрисовщиков: первая редакция писала сюда
    `= OPERATING_PROFIT + …`, и openpyxl клал такую клетку типом `f`. Excel
    пытается вычислить `OPERATING_PROFIT` как имя листа и печатает `#NAME?` —
    то есть источник расчёта был не просто сырым, а нечитаемым в самой книге.
    """
    assert _source_column(ROWS)["EBITDA"].data_type == "s"


def test_a_zero_substituted_for_a_missing_input_is_visible():
    rows = [*ROWS[:2], _calculated([
        {"enum": "OPERATING_PROFIT", "value": 28355.0, "substituted_zero": False},
        {"enum": "DEPRECIATION_AMORTIZATION", "value": 0.0, "substituted_zero": True},
    ])]
    assert "подставлен ноль" in _source_cells(rows)["EBITDA"]

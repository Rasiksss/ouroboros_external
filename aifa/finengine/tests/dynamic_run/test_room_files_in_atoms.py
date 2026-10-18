"""К16.3: файл из `out/` — результат узла временного атома: сохранён у владельца, с родословной и описанием (часть 16)."""
from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from finengine.config import settings
from finengine.contracts import DynamicArtifact, FileArtifact
from finengine.skills import sources
from finengine.skills.spec.ephemeral import AtomRequest, EphemeralAtom, to_atom_def
from finengine.skills.spec.graph_executor import RunInputs
from finengine.skills.spec.room_files import describe_file

OWNER, PROJECT = 7, "chat-1"

BOOK_CODE = '''
def run(inputs):
    import os
    from openpyxl import Workbook
    os.makedirs("out", exist_ok=True)
    wb = Workbook()
    data = wb.active
    data.title = "Данные"
    data.append(["Показатель", "Значение"])
    data.append(["Выручка", 100])
    data.append(["Себестоимость", 40])
    calc = wb.create_sheet("Расчёт")
    calc["A1"] = "Валовая прибыль"
    calc["B1"] = "=Данные!B2-Данные!B3"
    wb.save("out/итоги.xlsx")
    return {"filename": "итоги.xlsx", "sheets": 2}
'''


@pytest.fixture(autouse=True)
def _local_blobs(tmp_path, monkeypatch, catalog):  # оригиналы файлов — в подставном каталоге ouroboros (шаг 11)
    monkeypatch.setattr(settings, "s3_endpoint", "")
    monkeypatch.setattr(settings, "data_dir", tmp_path)


def _atom(code: str) -> EphemeralAtom:
    request = AtomRequest(
        name="eph.build_book", purpose="книга", output_type="dyn:book",
        output_schema={"type": "object"}, requires=(),
    )
    return EphemeralAtom(request, code)


def _ctx(**extra) -> SimpleNamespace:
    return SimpleNamespace(
        run_id="run-1", inputs=RunInputs(extra=extra), dependencies=lambda: (),
    )


def _run(code: str, **extra):
    extra = extra or {"owner_id": str(OWNER), "project_id": PROJECT}
    return to_atom_def(_atom(code)).executor(None, _ctx(**extra))


def _zip(files: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, text in files.items():
            archive.writestr(name, text)
    return buffer.getvalue()


# ───────────────────────── describe_file ─────────────────────────

def test_an_xlsx_is_described_by_its_sheets_and_formulas() -> None:
    workbook = Workbook()
    workbook.active.title = "Данные"
    workbook.active["A1"] = 1
    calc = workbook.create_sheet("Расчёт")
    calc["A1"] = "=Данные!A1*2"
    calc["A2"] = "=A1+1"
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = describe_file("итоги.xlsx", buffer.getvalue())

    assert "листов 2" in text and "Данные" in text and "Расчёт" in text and "формул 2" in text


def test_a_pptx_is_described_by_its_slides_and_a_docx_by_its_paragraphs() -> None:
    deck = _zip({f"ppt/slides/slide{i}.xml": "<p/>" for i in (1, 2, 3)} | {"ppt/presentation.xml": "<p/>"})
    doc = _zip({"word/document.xml": "<w:body><w:p>a</w:p><w:p>b</w:p></w:body>"})

    assert "слайдов 3" in describe_file("deck.pptx", deck)
    assert "абзацев 2" in describe_file("doc.docx", doc)


def test_an_unknown_format_is_described_by_size_and_hash_only() -> None:
    text = describe_file("схема.png", b"\x89PNG" + b"0" * 10)

    assert "png" in text and "14" in text


def test_a_broken_container_is_described_without_failing() -> None:
    assert "xlsx" in describe_file("битый.xlsx", b"PK\x03\x04not-a-zip")


# ───────────────────────── файл как результат узла ─────────────────────────

def test_the_file_written_to_out_becomes_a_file_artifact_with_lineage_to_the_nodes_result(local_room) -> None:
    results = _run(BOOK_CODE)

    dynamic = next(item for item in results if isinstance(item, DynamicArtifact))
    artifacts = [item for item in results if isinstance(item, FileArtifact)]
    assert dynamic.payload == {"filename": "итоги.xlsx", "sheets": 2}
    assert len(artifacts) == 1
    artifact = artifacts[0]
    assert artifact.file_kind == "xlsx" and artifact.filename == "итоги.xlsx"
    assert artifact.source_result.result_type == "DynamicArtifact"
    assert artifact.source_result.result_id == dynamic.evidence_id
    assert artifact.template_id == "eph.build_book" and artifact.template_version
    assert "листов 2" in artifact.summary and "формул 1" in artifact.summary


def test_the_stored_file_belongs_to_the_real_owner_and_is_readable_back(local_room) -> None:
    artifact = next(item for item in _run(BOOK_CODE) if isinstance(item, FileArtifact))

    file_id = artifact.object_store_location.rsplit("/", 1)[1].split(".")[0]
    content, meta = sources.read_source_bytes(OWNER, PROJECT, file_id)

    assert meta.filename == "итоги.xlsx"
    assert len(content) == artifact.size_bytes
    assert zipfile.is_zipfile(io.BytesIO(content))
    assert sources.read_source_bytes("system", PROJECT, file_id) is None


def test_a_node_without_files_returns_only_its_result(local_room) -> None:
    results = _run("def run(inputs):\n    return {'ok': True}\n")

    assert [type(item) for item in results] == [DynamicArtifact]


def test_a_foreign_format_under_an_office_extension_is_refused(local_room) -> None:
    code = 'def run(inputs):\n    import os\n    os.makedirs("out", exist_ok=True)\n    open("out/a.xlsx", "w").write("привет")\n    return {}\n'

    with pytest.raises(ValueError, match="a.xlsx"):
        _run(code)


def test_an_executable_in_out_is_refused(local_room) -> None:
    code = 'def run(inputs):\n    import os\n    os.makedirs("out", exist_ok=True)\n    open("out/run.exe", "wb").write(b"MZ")\n    return {}\n'

    with pytest.raises(ValueError, match="run.exe"):
        _run(code)


def test_a_file_without_a_known_owner_is_refused_instead_of_going_to_the_system_owner(local_room) -> None:
    with pytest.raises(ValueError, match="владел"):
        _run(BOOK_CODE, project_id=PROJECT)


def test_the_atom_code_still_cannot_read_the_report_text(local_room) -> None:
    ctx = SimpleNamespace(
        run_id="run-1", inputs=RunInputs(document_text="СЕКРЕТ ОТЧЁТА", extra={"owner_id": "7", "project_id": PROJECT}),
        dependencies=lambda: (),
    )
    code = "def run(inputs):\n    return {'seen': list(inputs)}\n"

    result = to_atom_def(_atom(code)).executor(None, ctx)

    assert result[0].payload == {"seen": []}

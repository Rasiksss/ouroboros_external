from __future__ import annotations

import hashlib
import io

import openpyxl
import pymupdf

from finengine.contracts import EvidenceRef, FormOutput, FormRow, OutputRunMetadata, canonical_result_hash
from finengine.skills.atoms.out_file import build_file_artifact, render_xlsx


def _form() -> FormOutput:
    return FormOutput(
        output_id="form-1",
        run_id="run-1",
        role="dgr",
        layout_id="dgr.bank",
        layout_version="1",
        preserve_empty=True,
        rows=[
            FormRow(
                row_id="dgr.bank:000", order=0, kind="dgr_row", enum="REVENUE",
                label="Выручка", value=125.0, display_value="125,0", status="supported",
                evidence_refs=[EvidenceRef(evidence_id="ev-1", document_version="doc-v2")],
            ),
            FormRow(
                row_id="dgr.bank:001", order=1, kind="dgr_row", enum="RWA",
                label="RWA", status="missing", missing_reason="evidence_not_found",
            ),
        ],
        run_metadata=OutputRunMetadata(run_id="run-1", evidence_version="evidence-v3"),
    )


def test_xlsx_is_a_pure_projection_and_artifact_records_lineage():
    source = _form()
    content = render_xlsx(source)
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    cells = [cell for row in workbook.active.iter_rows() for cell in row]

    assert not any(cell.data_type == "f" for cell in cells), "exporter must not calculate via XLSX formulas"
    assert workbook.active.cell(row=2, column=6).value == 125.0
    assert workbook.active.cell(row=2, column=7).value == "125,0"
    assert workbook.active.cell(row=3, column=5).value == "missing"

    artifact = build_file_artifact(
        source=source,
        file_kind="xlsx",
        template_id="dgr-xlsx",
        template_version="2.4",
        location="objects/run-1/form-1.xlsx",
        content=content,
        filename="dgr.xlsx",
    )
    assert artifact.content_hash == hashlib.sha256(content).hexdigest()
    assert artifact.source_result.result_id == source.output_id
    assert artifact.source_result.result_hash == canonical_result_hash(source)
    assert artifact.source_result.evidence_version == "evidence-v3"
    assert artifact.template_version == "2.4"


def test_repeated_render_keeps_values_and_hash_stable_and_pdf_opens():
    source = _form()
    first = render_xlsx(source)
    second = render_xlsx(source)
    assert first == second
    assert hashlib.sha256(first).digest() == hashlib.sha256(second).digest()

    from finengine.skills.atoms.out_file import render_pdf
    pdf = render_pdf(source)
    document = pymupdf.open(stream=pdf, filetype="pdf")
    try:
        assert document.page_count >= 1
        text = "".join(page.get_text() for page in document)
        assert "Выручка" in text
        assert "125,0" in text
    finally:
        document.close()


def test_the_artifact_of_the_standard_exporter_describes_its_structure_for_the_planner():
    """Часть 16: без описания планировщик цикла отдавал значения-таблицу как «Excel с формулами на двух листах»."""
    source = _form()

    artifact = build_file_artifact(
        source=source, file_kind="xlsx", template_id="FormOutput.xlsx", template_version="1.0",
        location="objects/run-1/form-1.xlsx", content=render_xlsx(source), filename="form.xlsx",
    )

    assert "листов 1" in artifact.summary and "формул 0" in artifact.summary

"""Сторожа графика «от операционного потока к свободному» (#157).

ЧТО ОХРАНЯЮТ. Эталонный отчёт, с которого списан этот график, НЕПРОВЕРЯЕМ: EBITDA
в его HTML — 41 150, а в PDF того же прогона — 32 393, расхождение 27% между
двумя артефактами одного расчёта. Так бывает, когда картинка считает своё, а
таблица — своё. Здесь проверяется ровно обратное: график берёт числа из ТАБЛИЦЫ
ЗАЗЕМЛЕНИЯ, сходится с ней, а если не сходится — не рисуется и называет причину.

ПУТЬ ТОТ ЖЕ, ЧТО У ПОЛЬЗОВАТЕЛЯ. Спецификация собирается `build_charts` — той
самой функцией, которую зовёт `llm/tools.py` перед `parse_report`, — и уезжает
в карточку тем же вызовом `parse_report(..., charts=...)`. Лист Excel собирается
настоящим `build_metrics_xlsx` и вычитывается обратно openpyxl'ом. Ни одного
внутреннего помощника в обход живого пути.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

# src-layout: делаем `finengine` импортируемым без editable-установки (как в
# соседних тестах карточки).
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from finengine.skills.analysis.charts import build_charts, waterfall_ocf_to_fcf  # noqa: E402
from finengine.skills.analysis.parse import parse_report  # noqa: E402
from finengine.skills.metrics_xlsx import build_metrics_xlsx  # noqa: E402
from finengine.skills.reasoner.runner import build_metrics_table_str  # noqa: E402

UNIT = "млн. руб."

# Числа сходятся так же, как их сводит движок: FREE_CASH_FLOW считает КОД по
# формуле реестра `NET_OPERATING_CASH_FLOW − CAPEX`, причём CAPEX входит в
# формулу по модулю (`formula_default_abs_behavior=true`). 120 000 − 85 000.
_OCF = 120_000.0
_CAPEX = 85_000.0
_FCF = 35_000.0
_INTEREST = 21_000.0
# Аренда взята такой, чтобы остаток ушёл В МИНУС (35 000 − 21 000 − 19 000):
# это и есть тот случай, ради которого график заведён, — «бизнес не окупает
# собственный долг». Столбик обязан уйти ниже нуля, а не упереться в него.
_LEASE = 19_000.0


def _metric(enum: str, name_ru: str, value: float | None, unit: str = UNIT) -> dict:
    """Событие движка в том виде, в каком оно приходит на шину заземления."""
    return {
        "type": "metric", "enum": enum, "name_ru": name_ru,
        "value_current": value, "value_prev": None,
        "period_values": ([{"period": "FY 2024", "value": value}] if value is not None else []),
        "unit": unit, "found": value is not None,
    }


def grounding_metrics(**over: float | None) -> list[dict]:
    """Таблица заземления прогона: пять входов графика плюс посторонний
    показатель — график обязан выбирать нужные строки, а не первые пять."""
    v = {"NET_OPERATING_CASH_FLOW": _OCF, "CAPEX": _CAPEX, "FREE_CASH_FLOW": _FCF,
         "INTEREST_PAID": _INTEREST, "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES": _LEASE}
    v.update(over)
    names = {
        "NET_OPERATING_CASH_FLOW": "Чистый операционный денежный поток",
        "CAPEX": "Капитальные затраты",
        "FREE_CASH_FLOW": "Свободный денежный поток",
        "INTEREST_PAID": "Уплаченные проценты",
        "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES": "Погашение обязательств по аренде",
    }
    rows = [_metric(e, names[e], v[e]) for e in names]
    rows.append(_metric("SALES_REVENUE", "Выручка", 900_000.0))
    return rows


def _points(spec: dict) -> dict[str, float]:
    return {p["label"]: p["value"] for p in spec["points"]}


# ───────── Сторож 1: ИНВАРИАНТ — график и таблица показывают одно число ─────────


def test_chart_values_equal_grounding_table_values():
    """Каждый столбик = строка таблицы заземления, а не пересчёт по-своему.

    Сверяем со СТРОКОЙ ТАБЛИЦЫ, которую видит модель (`build_metrics_table_str`
    из промпта анализа), а не с нашей же копией чисел: разойдись график с
    промптом — вердикт и картинка будут про разные компании.
    """
    metrics = grounding_metrics()
    spec = build_charts(metrics)[0]
    table = build_metrics_table_str(metrics)

    by_label = _points(spec)
    # Столбики оттоков рисуются вниз (со знаком минус), в таблице они лежат
    # величиной — сверяем по модулю, как их берёт и калькулятор реестра.
    assert abs(by_label["Чистый операционный денежный поток"]) == _OCF
    assert abs(by_label["Капитальные затраты"]) == _CAPEX
    assert abs(by_label["Свободный денежный поток"]) == _FCF
    assert abs(by_label["Уплаченные проценты"]) == _INTEREST
    assert abs(by_label["Погашение обязательств по аренде"]) == _LEASE
    # И те же числа — в тексте таблицы, ушедшей в промпт.
    for value in (_OCF, _CAPEX, _FCF, _INTEREST, _LEASE):
        assert f"{value:,.0f}".replace(",", " ") in table


def test_free_cash_flow_is_read_from_table_not_recomputed():
    """Точка «свободный поток» — из таблицы. Разошлась с «OCF − capex» —
    графика нет и названа причина.

    Это и есть защита от дефекта эталона: там одно и то же число в двух его
    артефактах отличалось на 27%, потому что никто не сверял.
    """
    spec, reason = waterfall_ocf_to_fcf(grounding_metrics(FREE_CASH_FLOW=44_500.0))
    assert spec is None
    assert "не сходится" in reason
    assert "44 500" in reason and "35 000" in reason  # оба числа названы


def test_waterfall_steps_add_up_to_the_final_bar():
    """Итоговый столбик = сумма показанных столбиков. Водопад, который не
    сходится в собственный итог, врёт формой, а не цифрой."""
    spec = build_charts(grounding_metrics())[0]
    values = [p["value"] for p in spec["points"]]
    # points = [OCF, −capex, ИТОГ FCF, −проценты, −аренда, ИТОГ остаток]
    assert values[0] + values[1] == pytest.approx(values[2])
    assert values[2] + values[3] + values[4] == pytest.approx(values[5])
    assert values[5] < 0  # смысл графика: потока не хватает на долг и аренду


# ───────── Сторож 2: масштаб оси выводится ИЗ ДАННЫХ ─────────


def test_unit_comes_from_the_data():
    for unit in ("млн. руб.", "тыс. руб.", "USD mn"):
        rows = [{**m, "unit": unit} for m in grounding_metrics()]
        assert build_charts(rows)[0]["unit"] == unit


def test_mixed_units_refuse_with_a_reason():
    """Тысячи и миллионы в одном водопаде — сложение несравнимого: все числа
    верные, картинка лживая. Отказ с причиной, а не молча другой масштаб."""
    rows = grounding_metrics()
    for m in rows:
        if m["enum"] == "CAPEX":
            m["unit"] = "тыс. руб."
    spec, reason = waterfall_ocf_to_fcf(rows)
    assert spec is None
    assert "единицы" in reason and "CAPEX" in reason


def test_missing_unit_refuses():
    """Безымянная единица = неизвестный масштаб. Рисовать «просто столбики»
    значит предложить читателю додумать масштаб самому."""
    rows = [{**m, "unit": ""} for m in grounding_metrics()]
    spec, reason = waterfall_ocf_to_fcf(rows)
    assert spec is None
    assert "единица" in reason


# ───────── Сторож 3: нет данных ≠ ноль ─────────


@pytest.mark.parametrize("enum", [
    "NET_OPERATING_CASH_FLOW", "CAPEX", "FREE_CASH_FLOW",
    "INTEREST_PAID", "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES",
])
def test_missing_input_refuses_instead_of_drawing_zero(enum: str):
    spec, reason = waterfall_ocf_to_fcf(grounding_metrics(**{enum: None}))
    assert spec is None
    assert enum in reason


def test_empty_grounding_gives_no_charts():
    """Анализ без прогона извлечения (станция заземлялась сама) — графика нет,
    карточка при этом собирается как раньше."""
    assert build_charts([]) == []


# ───────── Сторож 4: спецификация доезжает до карточки живым путём ─────────


def test_spec_reaches_the_card_through_parse_report():
    """Зеркало сборки из llm/tools.py: build_charts → parse_report(charts=...)."""
    charts = build_charts(grounding_metrics())
    card = parse_report("## Вердикт\nПоток отрицательный.\n",
                        company="Пример", period="", company_type="non_financial",
                        standard="МСФО", language="ru", charts=charts)
    assert len(card["charts"]) == 1
    spec = card["charts"][0]
    # Форма — ровно та, что объявлена типом AnalysisChart в web/src/lib/api.ts.
    assert spec["kind"] == "waterfall"
    assert spec["unit"] == UNIT
    assert {"label", "value", "kind"} <= set(spec["points"][0])
    assert {p["kind"] for p in spec["points"]} == {"start", "delta", "total"}
    # СПЕЦИФИКАЦИЯ, А НЕ КАРТИНКА: разметки в ней быть не должно — фронт строит
    # SVG элементами React, и любой тег отсюда означал бы, что кто-то завёл
    # второй путь рисования в обход запрета на dangerouslySetInnerHTML.
    assert "<" not in repr(spec)


# ───────── Сторож 5: число из графика проверяется руками в Excel ─────────


def test_excel_sheet_carries_the_same_numbers():
    import openpyxl

    metrics = grounding_metrics()
    wb = openpyxl.load_workbook(io.BytesIO(build_metrics_xlsx(metrics)))
    assert "Данные графиков" in wb.sheetnames
    ws = wb["Данные графиков"]
    rows = list(ws.values)
    header = list(rows[0])
    body = {r[2]: r for r in rows[1:]}

    assert header == ["График", "Ед.", "Столбик", "Значение", "Тип", "Источник"]
    spec = build_charts(metrics)[0]
    for point in spec["points"]:
        row = body[point["label"]]
        assert row[3] == point["value"], point["label"]  # то же число, что на графике
        assert row[1] == spec["unit"]                    # и та же единица


def test_excel_has_no_chart_sheet_when_chart_was_refused():
    """Лист не появляется «пустой шапкой»: отсутствие данных должно читаться
    как отсутствие данных, а не как график без столбиков."""
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(build_metrics_xlsx(grounding_metrics(CAPEX=None))))
    assert "Данные графиков" not in wb.sheetnames


# ── БОЕВАЯ ФОРМА ЕДИНИЦ ────────────────────────────────────────────────────
#
# Главный блокер, найденный состязательной проверкой: единицы у извлечённых и у
# рассчитанных показателей формируются РАЗНЫМИ ветками движка и строкой не
# совпадают. Пока сравнение было строковым, график отказывался на КАЖДОМ живом
# прогоне, и знал об этом только `log.info` — снаружи это выглядело как «график
# не сделали». Поэтому единицы здесь берутся не из головы, а у самого движка.


def _production_units():
    """Ровно те строки, что движок ставит на боевом пути."""
    from finengine.skills.bulk_extractor import _normalise_money_value, unit_for_calculated

    raw = "млн руб"  # типичный ответ `_detect_unit` по формам отчёта
    extracted = _normalise_money_value(120000, raw)[1]
    calculated = unit_for_calculated("NET_OPERATING_CASH_FLOW – CAPEX", raw)
    return extracted, calculated


def test_units_of_the_real_engine_do_differ_as_strings():
    """Предусловие блокера — фиксируем его, чтобы правка не «починила» тест
    вместо кода: если ветки когда-нибудь сойдутся, этот тест скажет об этом."""
    extracted, calculated = _production_units()
    assert extracted != calculated, (
        "ветки единиц сошлись — проверка ниже потеряла смысл, перечитать её")


def test_the_chart_is_built_on_units_as_the_engine_actually_writes_them():
    """ГЛАВНЫЙ: график обязан появиться на боевых данных, а не на удобных."""
    extracted, calculated = _production_units()
    spec, why = waterfall_ocf_to_fcf([
        {"enum": "NET_OPERATING_CASH_FLOW", "value_current": 120000, "unit": extracted},
        {"enum": "CAPEX", "value_current": -40000, "unit": extracted},
        {"enum": "INTEREST_PAID", "value_current": -10000, "unit": extracted},
        {"enum": "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES",
         "value_current": -5000, "unit": extracted},
        {"enum": "FREE_CASH_FLOW", "value_current": 80000, "unit": calculated},
    ])
    assert spec is not None, f"на боевых единицах график не строится: {why}"
    assert spec["unit"]


def test_a_single_blank_unit_refuses_instead_of_borrowing_the_neighbours_scale():
    """Пустая единица у ОДНОГО входа раньше молча отбрасывалась, и вход
    наследовал масштаб соседей — график с чужим ярлыком оси. Это ровно та
    ловушка, ради которой единица и берётся из данных: картинка лжёт, не соврав
    ни одним числом."""
    rows = [
        {"enum": "NET_OPERATING_CASH_FLOW", "value_current": 120000, "unit": "тыс. руб."},
        {"enum": "CAPEX", "value_current": -40000, "unit": "тыс. руб."},
        {"enum": "INTEREST_PAID", "value_current": -10000, "unit": ""},
        {"enum": "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES",
         "value_current": -5000, "unit": "тыс. руб."},
        {"enum": "FREE_CASH_FLOW", "value_current": 80000, "unit": "тыс. руб."},
    ]
    spec, why = waterfall_ocf_to_fcf(rows)
    assert spec is None, "вход без единицы занял масштаб у соседей"
    assert "INTEREST_PAID" in why, f"причина не называет виновника: {why}"


def test_genuinely_different_scales_are_still_refused():
    """Противовес: сравнение «по смыслу» не должно превратиться в «всё равно»."""
    spec, why = waterfall_ocf_to_fcf([
        {"enum": "NET_OPERATING_CASH_FLOW", "value_current": 120000, "unit": "млн. руб."},
        {"enum": "CAPEX", "value_current": -40000, "unit": "тыс. руб."},
        {"enum": "INTEREST_PAID", "value_current": -10000, "unit": "млн руб"},
        {"enum": "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES",
         "value_current": -5000, "unit": "млн руб"},
        {"enum": "FREE_CASH_FLOW", "value_current": 80000, "unit": "млн руб"},
    ])
    assert spec is None and "расходятся" in why


def test_an_unknown_unit_refuses_rather_than_assuming_roubles():
    spec, why = waterfall_ocf_to_fcf([
        {"enum": e, "value_current": v, "unit": "штук"}
        for e, v in (("NET_OPERATING_CASH_FLOW", 120000), ("CAPEX", -40000),
                     ("INTEREST_PAID", -10000),
                     ("PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES", -5000),
                     ("FREE_CASH_FLOW", 80000))
    ])
    assert spec is None and "не распознана" in why


# ── График не смеет ронять книгу извлечения ────────────────────────────────


def test_a_non_numeric_metric_does_not_destroy_the_extraction_workbook():
    """Регрессия, доказанная проверяющим исполнением: `build_charts` в сборке
    книги стоял БЕЗ обёртки, а «Немодифицированное» в аудиторском мнении —
    штатное значение реестра. Один такой показатель ронял `float()` внутри
    редьюсера и уносил ВСЮ книгу — единственный выход прогона, который
    пользователь забирает себе."""
    from finengine.skills.metrics_xlsx import build_metrics_xlsx

    data = build_metrics_xlsx([
        {"enum": "AUDIT_OPINION", "name_ru": "Аудиторское мнение",
         "value_current": "Немодифицированное", "unit": ""},
        {"enum": "SALES_REVENUE", "name_ru": "Выручка",
         "value_current": 100, "unit": "млн. руб."},
    ])
    assert isinstance(data, bytes) and len(data) > 1000, "книга не собралась"


# ── ДВА ЧИСЛА ПОД ОДНИМ ИМЕНЕМ: слышно, и вина названа верно (#163) ────────


def _duplicated_capex_stream():
    """Поток событий, который движок отдаёт РЕАЛЬНО.

    CAPEX заведён и извлекаемым, и расчётным (7 ячеек из 8), поэтому движок
    эмитит на него ДВА события: сначала извлечённое, потом посчитанное. Все
    прежние фикстуры выдавали по одному событию на enum — то есть охраняли
    данные, которых на проде не бывает.
    """
    u = "млн. руб."
    return [
        {"enum": "NET_OPERATING_CASH_FLOW", "name_ru": "Операционный поток",
         "value_current": 120000, "unit": u},
        {"enum": "CAPEX", "name_ru": "Капитальные затраты",
         "value_current": 85000, "unit": u},      # извлечён из отчёта
        {"enum": "INTEREST_PAID", "name_ru": "Проценты уплаченные",
         "value_current": -10000, "unit": u},
        {"enum": "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES", "name_ru": "Погашение аренды",
         "value_current": -5000, "unit": u},
        {"enum": "FREE_CASH_FLOW", "name_ru": "Свободный денежный поток",
         "value_current": 35000, "unit": u},      # посчитан от 85 000
        {"enum": "CAPEX", "name_ru": "Капитальные затраты",
         "value_current": 100000, "unit": u},     # ПОСЧИТАН — и затирает первый
    ]


def test_a_duplicated_metric_is_heard_not_silently_overwritten(caplog):
    """Молчание — главный дефект: второе событие затирало первое БЕЗ следа, и
    расхождение всплыло только состязательной проверкой."""
    import logging
    from finengine.skills.calc import dataset_from_metrics

    with caplog.at_level(logging.WARNING, logger="finengine.skills.calc"):
        table = dataset_from_metrics(_duplicated_capex_stream())

    assert any("#163" in r.message and "CAPEX" in r.message for r in caplog.records), (
        "дубль затёрся молча — ровно то, из-за чего расхождение жило незамеченным")
    assert table["CAPEX"].get("conflict") == [85000.0], (
        "след о расхождении не оставлен — отказу нечем будет объяснить причину")


def test_our_own_defect_is_not_blamed_on_the_report():
    """Читатель по формулировке решает, ЧЕМУ не верить. Сказать «отчёт не
    сходится» там, где у показателя два прочтения ВНУТРИ НАС, — это клевета на
    отчётность: числа в ней верны, а сверить их между собой нечем нам."""
    spec, why = waterfall_ocf_to_fcf(_duplicated_capex_stream())

    assert spec is None, "график построен вопреки расхождению"
    assert "НАШЕЙ причине" in why, f"вина переложена на отчёт: {why}"
    assert "Капитальные затраты" in why, "не названо, какой показатель разошёлся"
    assert "не сходится с" not in why, "осталась формулировка, винящая отчёт"


def test_a_real_report_discrepancy_is_still_blamed_on_the_report():
    """Противовес: починив клевету, нельзя разучиться называть настоящий дефект
    отчётности."""
    u = "млн. руб."
    spec, why = waterfall_ocf_to_fcf([
        {"enum": "NET_OPERATING_CASH_FLOW", "value_current": 120000, "unit": u},
        {"enum": "CAPEX", "value_current": 40000, "unit": u},
        {"enum": "INTEREST_PAID", "value_current": -10000, "unit": u},
        {"enum": "PRINCIPAL_REPAYMENT_OF_LEASE_LIABILITIES", "value_current": -5000, "unit": u},
        {"enum": "FREE_CASH_FLOW", "value_current": 99999, "unit": u},  # не бьётся
    ])
    assert spec is None
    assert "не сходится с" in why and "НАШЕЙ причине" not in why, why


def test_identical_duplicates_stay_quiet(caplog):
    """Дубль, совпавший до копейки, безвреден — шуметь о нём значит приучить
    читать логи по диагонали, а тогда и настоящее расхождение проглядят."""
    import logging
    from finengine.skills.calc import dataset_from_metrics

    rows = [{"enum": "EBITDA", "value_current": 500, "unit": "млн. руб."},
            {"enum": "EBITDA", "value_current": 500, "unit": "млн. руб."}]
    with caplog.at_level(logging.WARNING, logger="finengine.skills.calc"):
        table = dataset_from_metrics(rows)

    assert not [r for r in caplog.records if "#163" in r.message], (
        "поднят шум на безвредном дубле")
    assert "conflict" not in table["EBITDA"], "оставлен ложный след расхождения"

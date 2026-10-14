"""Real, minimal OOXML documents: what the attachment validator and parsers accept without stubs."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def write_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "Итоги"
    workbook.save(str(path))


def write_document(path: Path) -> None:
    document = Document()
    document.add_paragraph("Итоги")
    document.save(str(path))


def write_deck(path: Path) -> None:
    """One slide with a title, a bullet, a 2x2 table and speaker notes — every part the parser reports."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Итоги"
    slide.placeholders[1].text_frame.text = "первый пункт"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    for row, values in enumerate((("Регион", "Сумма"), ("Москва", "10"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    slide.notes_slide.notes_text_frame.text = "сказать вслух"
    presentation.save(str(path))

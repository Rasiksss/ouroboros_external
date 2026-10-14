from __future__ import annotations

import datetime
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from ouroboros.attachment_parser import _MAX_SLIDES, _parse_pptx, _validate_ooxml_package
from ouroboros.storage.attachment_validation import AttachmentError, validate_attachment_path


def _ooxml(path: Path, *, member: str, payload: bytes = b"<root/>") -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("_rels/.rels", b"<Relationships/>")
        archive.writestr(member, payload)


def test_ooxml_preflight_accepts_expected_xlsx_members(tmp_path):
    path = tmp_path / "book.xlsx"
    _ooxml(path, member="xl/workbook.xml")

    _validate_ooxml_package(
        path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_ooxml_preflight_rejects_extreme_compression_ratio(tmp_path):
    path = tmp_path / "book.xlsx"
    _ooxml(path, member="xl/workbook.xml", payload=b"0" * (2 * 1024 * 1024))

    with pytest.raises(ValueError, match="compression ratio"):
        _validate_ooxml_package(
            path,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


def test_ooxml_preflight_rejects_parent_traversal(tmp_path):
    path = tmp_path / "document.docx"
    _ooxml(path, member="word/document.xml")
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("../escape", b"bad")

    with pytest.raises(ValueError, match="unsafe path"):
        _validate_ooxml_package(
            path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


_PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Итоги"
    slide.placeholders[1].text_frame.text = "первый пункт"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    for row, values in enumerate((("Регион", "Сумма"), ("Москва", "10"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    slide.notes_slide.notes_text_frame.text = "сказать вслух"
    presentation.save(str(path))


def test_ooxml_preflight_accepts_expected_pptx_members(tmp_path):
    path = tmp_path / "deck.pptx"
    _ooxml(path, member="ppt/presentation.xml")

    _validate_ooxml_package(path, _PPTX_MIME)


def test_ooxml_preflight_rejects_a_pptx_without_its_presentation_part(tmp_path):
    path = tmp_path / "deck.pptx"
    _ooxml(path, member="word/document.xml")

    with pytest.raises(ValueError, match="missing required members"):
        _validate_ooxml_package(path, _PPTX_MIME)


def test_attachment_validation_accepts_a_pptx_deck(tmp_path):
    path = tmp_path / "deck.pptx"
    _deck(path)

    validated = validate_attachment_path(path, path.name, _PPTX_MIME)

    assert (validated.mime, validated.kind) == (_PPTX_MIME, "file")


def test_attachment_validation_rejects_a_pptx_that_is_not_a_package(tmp_path):
    path = tmp_path / "deck.pptx"
    path.write_bytes(b"plain text pretending to be a deck")

    with pytest.raises(AttachmentError, match="not a valid OOXML package"):
        validate_attachment_path(path, path.name, _PPTX_MIME)


def test_parsing_a_pptx_descends_into_grouped_shapes(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "снаружи"
    first = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    first.text_frame.text = "внутри группы"
    second = slide.shapes.add_textbox(Inches(4), Inches(3), Inches(2), Inches(1))
    second.text_frame.text = "рядом в группе"
    slide.shapes.add_group_shape([first, second])
    path = tmp_path / "grouped.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == [
        "## Slide 1",
        "снаружи",
        "внутри группы",
        "рядом в группе",
    ]


def test_parsing_a_pptx_survives_a_notes_part_without_a_body_placeholder(tmp_path):
    """has_notes_slide only promises the part exists; notes_text_frame answers None without a body placeholder."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Заголовок"
    placeholder = slide.notes_slide.notes_placeholder
    placeholder._element.getparent().remove(placeholder._element)
    path = tmp_path / "headless-notes.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == ["## Slide 1", "Заголовок"]


def test_parsing_a_pptx_reports_the_numbers_behind_a_chart(tmp_path):
    """A slide whose figures live in a chart would otherwise reach the model as a bare title."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Продажи"
    data = CategoryChartData()
    data.categories = ["Москва", "Питер"]
    data.add_series("2026", (12.0, 5.0))
    data.add_series("2027", (18.0, None))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(8), Inches(4), data)
    path = tmp_path / "chart.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == [
        "## Slide 1",
        "Продажи",
        "Chart:\tМосква\tПитер",
        "2026\t12.0\t5.0",
        "2027\t18.0\t",
    ]


def test_parsing_a_pptx_says_how_many_slides_it_left_out(tmp_path):
    """A silently shortened deck reads as complete, and the model answers about slides it never saw."""
    presentation = Presentation()
    for index in range(_MAX_SLIDES + 3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = f"слайд {index}"
    path = tmp_path / "long.pptx"
    presentation.save(str(path))

    extracted = _parse_pptx(path, "extract").split("\n")

    assert extracted[-3:] == [f"## Slide {_MAX_SLIDES}", f"слайд {_MAX_SLIDES - 1}", "[3 more slides omitted]"]


def test_parsing_a_pptx_keeps_a_multi_line_table_cell_on_its_own_row(tmp_path):
    """A cell split across lines would otherwise shift the row and pair the wrong value with the wrong column."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
    table.cell(0, 0).text = "Регион"
    table.cell(0, 1).text = "Сумма"
    table.cell(1, 0).text_frame.text = "Москва"
    table.cell(1, 0).text_frame.add_paragraph().text = "и область"
    table.cell(1, 1).text = "10"
    path = tmp_path / "wrapped-cell.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == [
        "## Slide 1",
        "Регион\tСумма",
        "Москва и область\t10",
    ]


def test_parsing_a_pptx_prefixes_every_line_of_multi_line_notes(tmp_path):
    """Without the prefix a continuation line is indistinguishable from text shown on the slide."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "тело слайда"
    frame = slide.notes_slide.notes_text_frame
    frame.text = "первая строка"
    frame.add_paragraph().text = "вторая строка"
    path = tmp_path / "multiline-notes.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == [
        "## Slide 1",
        "тело слайда",
        "Notes: первая строка",
        "Notes: вторая строка",
    ]


def test_parsing_a_pptx_decodes_a_date_axis(tmp_path):
    """A date axis caches Excel serial numbers, which carry no meaning for a reader."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = [datetime.date(2026, 1, 1), datetime.date(2026, 2, 1)]
    data.add_series("выручка", (10.0, 12.0))
    slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(2), Inches(8), Inches(4), data)
    path = tmp_path / "dated-chart.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == [
        "## Slide 1",
        "Chart:\t2026-01-01\t2026-02-01",
        "выручка\t10.0\t12.0",
    ]


def test_parsing_a_pptx_omits_the_header_row_of_a_chart_without_categories(tmp_path):
    """A scatter plot has no categories, and an empty header row tells the reader nothing."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = XyChartData()
    series = data.add_series("точки")
    series.add_data_point(1.0, 2.0)
    series.add_data_point(3.0, 4.0)
    slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(2), Inches(8), Inches(4), data)
    path = tmp_path / "scatter.pptx"
    presentation.save(str(path))

    assert _parse_pptx(path, "extract").split("\n") == ["## Slide 1", "точки\t2.0\t4.0"]


def test_parsing_a_pptx_turns_a_soft_line_break_into_a_newline(tmp_path):
    """python-pptx renders <a:br> as U+000B; the sibling parsers emit a newline for a soft break."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
    paragraph = frame.paragraphs[0]
    paragraph.add_run().text = "первая"
    paragraph._p.append(paragraph._p.makeelement(qn("a:br"), {}))
    paragraph.add_run().text = "вторая"
    path = tmp_path / "softbreak.pptx"
    presentation.save(str(path))

    extracted = _parse_pptx(path, "extract")

    assert "\v" not in extracted
    assert extracted.split("\n") == ["## Slide 1", "первая", "вторая"]


def test_parsing_a_pptx_reports_slides_tables_and_speaker_notes(tmp_path):
    path = tmp_path / "deck.pptx"
    _deck(path)

    assert _parse_pptx(path, "validate") == ""
    extracted = _parse_pptx(path, "extract").split("\n")

    assert extracted == [
        "## Slide 1",
        "Итоги",
        "первый пункт",
        "Регион\tСумма",
        "Москва\t10",
        "Notes: сказать вслух",
    ]

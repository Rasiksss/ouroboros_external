"""Document parsing in a spawned child without application dependencies."""

from __future__ import annotations

import datetime
import json
import os
import pathlib
import re
import sys
import time
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from typing import Any, Literal

from security_config.attachments import DOCX_MIME, PPTX_MIME, XLSX_MIME

MAX_EXTRACTED_CHARS = 100_000
PARSER_MEMORY_BYTES = 512 * 1024 * 1024
_MAX_ARCHIVE_ENTRIES = 10_000
_MAX_ARCHIVE_ENTRY_BYTES = 50 * 1024 * 1024
_MAX_ARCHIVE_TOTAL_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_COMPRESSION_RATIO = 200
_TRUNCATION_MARKER = "\n[TRUNCATED]"

_MAX_SLIDES = 200
# Excel counts days from 1900-01-01 as day 1 and carries a phantom 1900-02-29, so the epoch is two days back.
_EXCEL_EPOCH = datetime.date(1899, 12, 30)
_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_FORMAT_CODE_PATH = f".//{{{_CHART_NS}}}formatCode"
# 'm' alone also means minutes, so a year or a day is what marks a format as a date.
_DATE_FORMAT_RE = re.compile(r"[yd]", re.IGNORECASE)


@dataclass(frozen=True)
class ParserResult:
    status: Literal["ok", "unavailable", "invalid"]
    payload: str
    started_at: float
    finished_at: float
    import_seconds: float | None
    parse_seconds: float | None


@dataclass
class _ImportTiming:
    seconds: float = 0.0


@contextmanager
def _measure_import(timing: _ImportTiming | None) -> Iterator[None]:
    started = time.monotonic()
    try:
        yield
    finally:
        if timing is not None:
            timing.seconds += time.monotonic() - started


def _decode_text(raw: bytes) -> str:
    encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp1251")
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            if "\x00" in text:
                raise ValueError("text file contains binary NUL bytes")
            return text
        except UnicodeDecodeError:
            continue
    raise ValueError("text encoding is not supported")


def _bounded_text(value: str) -> str:
    if len(value) <= MAX_EXTRACTED_CHARS:
        return value
    return value[: MAX_EXTRACTED_CHARS - len(_TRUNCATION_MARKER)] + _TRUNCATION_MARKER


def _validate_ooxml_package(path: pathlib.Path, mime: str) -> None:
    required = {
        XLSX_MIME: {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"},
        DOCX_MIME: {"[Content_Types].xml", "_rels/.rels", "word/document.xml"},
        PPTX_MIME: {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"},
    }.get(mime)
    if required is None:
        raise ValueError("unsupported OOXML type")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > _MAX_ARCHIVE_ENTRIES:
                raise ValueError("OOXML package has too many entries")
            names: set[str] = set()
            total_size = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                parts = pathlib.PurePosixPath(normalized).parts
                if normalized.startswith("/") or ".." in parts:
                    raise ValueError("OOXML package contains an unsafe path")
                if entry.flag_bits & 1:
                    raise ValueError("encrypted OOXML is not supported")
                names.add(normalized)
                if entry.is_dir():
                    continue
                if entry.file_size > _MAX_ARCHIVE_ENTRY_BYTES:
                    raise ValueError("OOXML entry exceeds size limit")
                total_size += entry.file_size
                if total_size > _MAX_ARCHIVE_TOTAL_BYTES:
                    raise ValueError("OOXML package exceeds size limit")
                if entry.file_size / max(1, entry.compress_size) > _MAX_ARCHIVE_COMPRESSION_RATIO:
                    raise ValueError("OOXML entry exceeds compression ratio limit")
            if not required.issubset(names):
                raise ValueError("OOXML package is missing required members")
    except zipfile.BadZipFile as exc:
        raise ValueError("file is not a valid OOXML package") from exc


def _parse_pdf(path: pathlib.Path, mode: str, timing: _ImportTiming | None = None) -> str:
    # Optional dependencies load after limits; measure their cold import separately.
    with _measure_import(timing):
        from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF is not supported")
    if mode == "validate":
        len(reader.pages)
        return ""
    pages: list[str] = []
    length = 0
    for index, page in enumerate(reader.pages[:100], start=1):
        value = f"## Page {index}\n{page.extract_text() or ''}"
        pages.append(value)
        length += len(value)
        if length >= MAX_EXTRACTED_CHARS:
            break
    return "\n\n".join(pages)


def _parse_xlsx(path: pathlib.Path, mode: str, timing: _ImportTiming | None = None) -> str:
    _validate_ooxml_package(path, XLSX_MIME)
    # Optional dependencies load after limits; measure their cold import separately.
    with _measure_import(timing):
        from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if mode == "validate":
            return ""
        lines: list[str] = []
        length = 0
        for sheet in workbook.worksheets[:10]:
            header = f"## Sheet: {sheet.title}"
            lines.append(header)
            length += len(header)
            for row in sheet.iter_rows(max_row=500, max_col=100, values_only=True):
                value = "\t".join("" if cell is None else str(cell) for cell in row)
                lines.append(value)
                length += len(value)
                if length >= MAX_EXTRACTED_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)
    finally:
        workbook.close()


def _parse_docx(path: pathlib.Path, mode: str, timing: _ImportTiming | None = None) -> str:
    _validate_ooxml_package(path, DOCX_MIME)
    # Optional dependencies load after limits; measure their cold import separately.
    with _measure_import(timing):
        from docx import Document

    document = Document(str(path))
    if mode == "validate":
        len(document.paragraphs)
        return ""
    lines: list[str] = []
    length = 0
    for paragraph in document.paragraphs:
        if not paragraph.text:
            continue
        lines.append(paragraph.text)
        length += len(paragraph.text)
        if length >= MAX_EXTRACTED_CHARS:
            return "\n".join(lines)
    for table in document.tables:
        for row in table.rows:
            value = "\t".join(cell.text for cell in row.cells)
            lines.append(value)
            length += len(value)
            if length >= MAX_EXTRACTED_CHARS:
                return "\n".join(lines)
    return "\n".join(lines)


def _chart_lines(chart: Any) -> list[str]:
    """Render a chart as its cached table: a row of categories, then one row per series."""
    lines: list[str] = []
    for plot in chart.plots:
        categories = _chart_categories(plot)
        # A scatter or bubble plot has no categories, and an empty header row says nothing.
        if categories:
            lines.append("Chart:\t" + "\t".join(categories))
        lines.extend(
            "\t".join([str(series.name), *("" if value is None else str(value) for value in series.values)])
            for series in plot.series
        )
    return lines


def _chart_categories(plot: Any) -> list[str]:
    """A date axis caches its labels as Excel serial numbers, which mean nothing to a reader."""
    labels = [str(category) for category in plot.categories]
    if not _DATE_FORMAT_RE.search(plot._element.findtext(_FORMAT_CODE_PATH) or ""):
        return labels
    decoded: list[str] = []
    for label in labels:
        try:
            decoded.append((_EXCEL_EPOCH + datetime.timedelta(days=float(label))).isoformat())
        except ValueError:
            decoded.append(label)
    return decoded


def _parse_pptx(path: pathlib.Path, mode: str, timing: _ImportTiming | None = None) -> str:
    _validate_ooxml_package(path, PPTX_MIME)
    # Optional dependencies load after limits; measure their cold import separately.
    with _measure_import(timing):
        from pptx import Presentation
        from pptx.shapes.group import GroupShape

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    if mode == "validate":
        return ""
    lines: list[str] = []
    length = 0
    for index, slide in enumerate(slides[:_MAX_SLIDES], start=1):
        values = [f"## Slide {index}"]
        # A group has neither a text frame nor a table, so its text is lost unless the walk descends.
        pending = list(slide.shapes)
        while pending:
            shape = pending.pop(0)
            if isinstance(shape, GroupShape):
                pending[:0] = list(shape.shapes)
            elif shape.has_text_frame:
                # A soft line break arrives as U+000B, and every sibling parser emits a newline for one.
                values.extend(p.text.replace("\v", "\n") for p in shape.text_frame.paragraphs if p.text)
            elif shape.has_table:
                # A cell may hold several paragraphs, and a row has to stay one line to keep its columns.
                values.extend("\t".join(" ".join(c.text.split()) for c in row.cells) for row in shape.table.rows)
            elif shape.has_chart:
                values.extend(_chart_lines(shape.chart))
        # A notes part without a body placeholder answers None, and has_notes_slide does not say so.
        frame = slide.notes_slide.notes_text_frame if slide.has_notes_slide else None
        # Every notes line keeps the prefix: without it a continuation reads as text shown on the slide.
        notes = (frame.text if frame is not None else "").replace("\v", "\n")
        values.extend(f"Notes: {line.strip()}" for line in notes.splitlines() if line.strip())
        for value in values:
            lines.append(value)
            length += len(value)
            if length >= MAX_EXTRACTED_CHARS:
                return "\n".join(lines)
    # Without this the deck reads as complete, and the model answers confidently about slides it never saw.
    if len(slides) > _MAX_SLIDES:
        lines.append(f"[{len(slides) - _MAX_SLIDES} more slides omitted]")
    return "\n".join(lines)


def _parse_path(path: pathlib.Path, mime: str, mode: str, timing: _ImportTiming | None = None) -> str:
    if mime in {"text/plain", "text/markdown", "text/csv", "application/json"}:
        text = _decode_text(path.read_bytes())
        if mime == "application/json":
            json.loads(text)
        return "" if mode == "validate" else text
    if mime == "application/pdf":
        return _parse_pdf(path, mode, timing)
    if mime == XLSX_MIME:
        return _parse_xlsx(path, mode, timing)
    if mime == DOCX_MIME:
        return _parse_docx(path, mode, timing)
    if mime == PPTX_MIME:
        return _parse_pptx(path, mode, timing)
    raise ValueError(f"unsupported parser MIME: {mime}")


def _apply_resource_limits(cpu_seconds: int) -> None:
    # Unix-only limits are loaded in the child, not by the portable tool facade.
    import resource

    try:
        resource.setrlimit(resource.RLIMIT_AS, (PARSER_MEMORY_BYTES, PARSER_MEMORY_BYTES))
    except ValueError:
        if sys.platform != "darwin":
            raise
    resource.setrlimit(resource.RLIMIT_CPU, (cpu_seconds, cpu_seconds))


def parser_child(connection: Any, path: str, mime: str, mode: str, cpu_seconds: int) -> None:
    started_at = time.monotonic()
    parse_started: float | None = None
    timing = _ImportTiming()
    status: Literal["ok", "unavailable", "invalid"] = "ok"
    payload = ""
    try:
        try:
            os.environ.update(OPENBLAS_NUM_THREADS="1", OMP_NUM_THREADS="1", MKL_NUM_THREADS="1")
            _apply_resource_limits(cpu_seconds)
            started_at = parse_started = time.monotonic()
            payload = _bounded_text(_parse_path(pathlib.Path(path), mime, mode, timing))
        except ModuleNotFoundError as exc:
            status, payload = "unavailable", str(exc.name or "parser dependency")
        except Exception as exc:
            status, payload = "invalid", str(exc)
        finished_at = time.monotonic()
        import_seconds = timing.seconds if parse_started is not None else None
        parse_seconds = finished_at - parse_started - timing.seconds if parse_started is not None else None
        connection.send(ParserResult(status, payload, started_at, finished_at, import_seconds, parse_seconds))
    finally:
        connection.close()

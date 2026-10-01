"""Explicit, isolated parsing for attachments bound to the current task."""

from __future__ import annotations

import json
import multiprocessing
import os
import pathlib
import re
import sys
import tempfile
import threading
import time
import zipfile
from dataclasses import dataclass
from typing import Any, Literal

from new_core.tools.context import ToolContext
from new_core.tools.entry import ToolEntry
from new_core.tools.schema import ToolSchema
from ouroboros.storage.blob.factory import create_storage_provider
from ouroboros.storage.blob.protocol import ObjectTooLargeError
from ouroboros.storage.blob.scopes import SCOPE_UPLOADS
from ouroboros.storage.chat_attachments import (
    MAX_DOCUMENT_BYTES,
    MAX_IMAGE_BYTES,
    AttachmentError,
    get_attachment,
)

MAX_EXTRACTED_CHARS = 100_000
PARSER_MEMORY_BYTES = 512 * 1024 * 1024
PARSER_CPU_SECONDS = 30
PARSER_WALL_TIMEOUT_SECONDS = 45.0
_MAX_ARCHIVE_ENTRIES = 10_000
_MAX_ARCHIVE_ENTRY_BYTES = 50 * 1024 * 1024
_MAX_ARCHIVE_TOTAL_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_COMPRESSION_RATIO = 200
_TRUNCATION_MARKER = "\n[TRUNCATED]"
_PARSER_SLOT = threading.BoundedSemaphore(1)

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class _FormatRule:
    mime: str
    aliases: frozenset[str]
    kind: Literal["image", "file"] = "file"


@dataclass(frozen=True)
class ValidatedAttachment:
    filename: str
    mime: str
    kind: Literal["image", "file"]
    size: int


_FORMATS = {
    ".pdf": _FormatRule("application/pdf", frozenset({"application/pdf"})),
    ".xlsx": _FormatRule(_XLSX_MIME, frozenset({_XLSX_MIME})),
    ".docx": _FormatRule(_DOCX_MIME, frozenset({_DOCX_MIME})),
    ".txt": _FormatRule("text/plain", frozenset({"text/plain"})),
    ".md": _FormatRule("text/markdown", frozenset({"text/markdown", "text/plain"})),
    ".csv": _FormatRule("text/csv", frozenset({"text/csv", "text/plain"})),
    ".json": _FormatRule("application/json", frozenset({"application/json"})),
    ".png": _FormatRule("image/png", frozenset({"image/png"}), "image"),
    ".jpg": _FormatRule("image/jpeg", frozenset({"image/jpeg"}), "image"),
    ".jpeg": _FormatRule("image/jpeg", frozenset({"image/jpeg"}), "image"),
    ".webp": _FormatRule("image/webp", frozenset({"image/webp"}), "image"),
    ".gif": _FormatRule("image/gif", frozenset({"image/gif"}), "image"),
}


def attachment_mime_for_suffix(filename: str) -> str:
    """Return the canonical MIME for an allowed filename suffix."""
    rule = _FORMATS.get(pathlib.Path(str(filename or "")).suffix.lower())
    return rule.mime if rule is not None else ""


def _current_attachment(ctx: ToolContext, attachment_id: str) -> dict[str, Any] | None:
    requested = str(attachment_id or "").strip()
    return next(
        (metadata for metadata in ctx.attachments if str(metadata.get("id") or "").strip() == requested),
        None,
    )


def _decode_text(raw: bytes) -> str:
    encodings = ("utf-16",) if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp1251")
    for encoding in encodings:
        try:
            text = raw.decode(encoding)
            if "\x00" in text:
                raise ValueError("text file contains binary NUL bytes")
            return text
        except UnicodeDecodeError:
            continue
    raise ValueError("text encoding is not supported")


def _bounded_text(value: str) -> str:
    if len(value) <= MAX_EXTRACTED_CHARS:
        return value
    return value[: MAX_EXTRACTED_CHARS - len(_TRUNCATION_MARKER)] + _TRUNCATION_MARKER


def _image_mime(path: pathlib.Path) -> str:
    with path.open("rb") as handle:
        head = handle.read(16)
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif"
    if len(head) >= 12 and head.startswith(b"RIFF") and head[8:12] == b"WEBP":
        return "image/webp"
    return ""


def _validate_ooxml_package(path: pathlib.Path, mime: str) -> None:
    required = {
        _XLSX_MIME: {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"},
        _DOCX_MIME: {"[Content_Types].xml", "_rels/.rels", "word/document.xml"},
    }.get(mime)
    if required is None:
        raise AttachmentError("invalid_attachment", "unsupported OOXML type", filename=path.name)
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > _MAX_ARCHIVE_ENTRIES:
                raise AttachmentError("invalid_attachment", "OOXML package has too many entries", filename=path.name)
            names: set[str] = set()
            total_size = 0
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                parts = pathlib.PurePosixPath(normalized).parts
                if normalized.startswith("/") or ".." in parts:
                    raise AttachmentError(
                        "invalid_attachment", "OOXML package contains an unsafe path", filename=path.name
                    )
                if entry.flag_bits & 0x1:
                    raise AttachmentError("invalid_attachment", "encrypted OOXML is not supported", filename=path.name)
                names.add(normalized)
                if entry.is_dir():
                    continue
                if entry.file_size > _MAX_ARCHIVE_ENTRY_BYTES:
                    raise AttachmentError("invalid_attachment", "OOXML entry exceeds size limit", filename=path.name)
                total_size += entry.file_size
                if total_size > _MAX_ARCHIVE_TOTAL_BYTES:
                    raise AttachmentError("invalid_attachment", "OOXML package exceeds size limit", filename=path.name)
                if entry.file_size / max(1, entry.compress_size) > _MAX_ARCHIVE_COMPRESSION_RATIO:
                    raise AttachmentError(
                        "invalid_attachment", "OOXML entry exceeds compression ratio limit", filename=path.name
                    )
            if not required.issubset(names):
                raise AttachmentError(
                    "invalid_attachment", "OOXML package is missing required members", filename=path.name
                )
    except zipfile.BadZipFile as exc:
        raise AttachmentError("invalid_attachment", "file is not a valid OOXML package", filename=path.name) from exc


def _parse_pdf(path: pathlib.Path, mode: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError("encrypted PDF is not supported")
    if mode == "validate":
        len(reader.pages)
        return ""
    pages: list[str] = []
    length = 0
    for index, page in enumerate(reader.pages[:100], start=1):
        value = f"## Page {index}\n{page.extract_text() or ''}"
        pages.append(value)
        length += len(value)
        if length >= MAX_EXTRACTED_CHARS:
            break
    return "\n\n".join(pages)


def _parse_xlsx(path: pathlib.Path, mode: str) -> str:
    _validate_ooxml_package(path, _XLSX_MIME)
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if mode == "validate":
            return ""
        lines: list[str] = []
        length = 0
        for sheet in workbook.worksheets[:10]:
            header = f"## Sheet: {sheet.title}"
            lines.append(header)
            length += len(header)
            for row in sheet.iter_rows(max_row=500, max_col=100, values_only=True):
                value = "\t".join("" if cell is None else str(cell) for cell in row)
                lines.append(value)
                length += len(value)
                if length >= MAX_EXTRACTED_CHARS:
                    return "\n".join(lines)
        return "\n".join(lines)
    finally:
        workbook.close()


def _parse_docx(path: pathlib.Path, mode: str) -> str:
    _validate_ooxml_package(path, _DOCX_MIME)
    from docx import Document

    document = Document(str(path))
    if mode == "validate":
        len(document.paragraphs)
        return ""
    lines: list[str] = []
    length = 0
    for paragraph in document.paragraphs:
        if not paragraph.text:
            continue
        lines.append(paragraph.text)
        length += len(paragraph.text)
        if length >= MAX_EXTRACTED_CHARS:
            return "\n".join(lines)
    for table in document.tables:
        for row in table.rows:
            value = "\t".join(cell.text for cell in row.cells)
            lines.append(value)
            length += len(value)
            if length >= MAX_EXTRACTED_CHARS:
                return "\n".join(lines)
    return "\n".join(lines)


def _parse_path(path: pathlib.Path, mime: str, mode: str) -> str:
    if mime in {"text/plain", "text/markdown", "text/csv", "application/json"}:
        text = _decode_text(path.read_bytes())
        if mime == "application/json":
            json.loads(text)
        return "" if mode == "validate" else text
    if mime == "application/pdf":
        return _parse_pdf(path, mode)
    if mime == _XLSX_MIME:
        return _parse_xlsx(path, mode)
    if mime == _DOCX_MIME:
        return _parse_docx(path, mode)
    raise ValueError(f"unsupported parser MIME: {mime}")


def _apply_resource_limits() -> None:
    import resource

    try:
        resource.setrlimit(resource.RLIMIT_AS, (PARSER_MEMORY_BYTES, PARSER_MEMORY_BYTES))
    except ValueError:
        if sys.platform != "darwin":
            raise
    resource.setrlimit(resource.RLIMIT_CPU, (PARSER_CPU_SECONDS, PARSER_CPU_SECONDS))


def _constrain_blas_threads() -> None:
    """Cap BLAS threads before numpy/OpenBLAS loads in the parser child.

    OpenBLAS reserves a per-thread buffer at startup and RLIMIT_AS counts that
    reservation, not the resident use: on a many-core node the child dies under
    the 512 MB address-space cap while importing openpyxl, which pulls in numpy.
    The parsers here are single-threaded XML/text readers, so one thread is plenty.
    """
    os.environ.update(
        {
            "OPENBLAS_NUM_THREADS": "1",
            "OMP_NUM_THREADS": "1",
            "MKL_NUM_THREADS": "1",
        }
    )


def _parser_child(connection: Any, path: str, mime: str, mode: str) -> None:
    try:
        _constrain_blas_threads()
        _apply_resource_limits()
        connection.send(("ok", _bounded_text(_parse_path(pathlib.Path(path), mime, mode))))
    except ModuleNotFoundError as exc:
        connection.send(("unavailable", str(exc.name or "parser dependency")))
    except Exception as exc:
        connection.send(("invalid", str(exc)))
    finally:
        connection.close()


def run_parser(
    path: pathlib.Path,
    mime: str,
    *,
    mode: Literal["validate", "extract"],
    timeout_sec: float = PARSER_WALL_TIMEOUT_SECONDS,
) -> str:
    timeout = max(0.0, float(timeout_sec))
    deadline = time.monotonic() + timeout
    if not _PARSER_SLOT.acquire(timeout=timeout):
        raise AttachmentError("invalid_attachment", "attachment parsing timed out", filename=path.name)
    try:
        process_context = multiprocessing.get_context("spawn")
        receive, send = process_context.Pipe(duplex=False)
        process = process_context.Process(
            target=_parser_child,
            args=(send, str(path), str(mime), mode),
            daemon=True,
        )
        process.start()
        send.close()
        try:
            if not receive.poll(max(0.0, deadline - time.monotonic())):
                process.terminate()
                process.join(timeout=1)
                if process.is_alive():
                    process.kill()
                    process.join(timeout=1)
                raise AttachmentError("invalid_attachment", "attachment parsing timed out", filename=path.name)
            status, payload = receive.recv()
        except EOFError as exc:
            raise AttachmentError(
                "invalid_attachment", "attachment parser exited unexpectedly", filename=path.name
            ) from exc
        finally:
            receive.close()
            if process.is_alive():
                process.join(timeout=1)
            if process.is_alive():
                process.terminate()
                process.join(timeout=1)
        if status == "ok":
            return str(payload)
        if status == "unavailable":
            raise AttachmentError(
                "attachment_parser_unavailable",
                f"attachment parser dependency is unavailable: {payload}",
                filename=path.name,
                retryable=True,
            )
        raise AttachmentError("invalid_attachment", str(payload), filename=path.name)
    finally:
        _PARSER_SLOT.release()


def validate_attachment_path(
    path: pathlib.Path,
    filename: str,
    declared_mime: str,
    *,
    image_max_bytes: int | None = None,
) -> ValidatedAttachment:
    clean_name = re.sub(
        r"[\x00-\x1f\x7f]",
        "_",
        pathlib.PurePosixPath(str(filename or "").replace("\\", "/")).name,
    )
    rule = _FORMATS.get(pathlib.Path(clean_name).suffix.lower())
    if rule is None:
        raise AttachmentError("attachment_type_not_allowed", "unsupported file format", filename=clean_name)
    mime = str(declared_mime or "").split(";", 1)[0].strip().lower()
    if mime not in rule.aliases:
        raise AttachmentError(
            "attachment_type_not_allowed",
            "file extension and MIME type do not match",
            filename=clean_name,
        )
    size = path.stat().st_size
    limit = (
        (MAX_IMAGE_BYTES if image_max_bytes is None else image_max_bytes)
        if rule.kind == "image"
        else MAX_DOCUMENT_BYTES
    )
    if size > limit:
        raise AttachmentError(
            "attachment_too_large",
            f"attachment exceeds {limit // (1024 * 1024)} MB limit",
            filename=clean_name,
        )
    if rule.kind == "image":
        if _image_mime(path) != rule.mime:
            raise AttachmentError("invalid_attachment", "file content does not match image type", filename=clean_name)
    else:
        with path.open("rb") as handle:
            head = handle.read(8)
        if rule.mime == "application/pdf" and not head.startswith(b"%PDF-"):
            raise AttachmentError("invalid_attachment", "file content does not match PDF", filename=clean_name)
        if rule.mime in {_XLSX_MIME, _DOCX_MIME} and not head.startswith(b"PK"):
            raise AttachmentError("invalid_attachment", "file is not a valid OOXML package", filename=clean_name)
        run_parser(path, rule.mime, mode="validate")
    return ValidatedAttachment(clean_name, rule.mime, rule.kind, size)


def _attachment_read(ctx: ToolContext, attachment_id: str) -> str:
    current = _current_attachment(ctx, attachment_id)
    if current is None:
        return "Attachment is not available in the current task."
    provider = create_storage_provider(data_dir=ctx.drive_root)
    metadata = get_attachment(provider, attachment_id)
    if metadata is None:
        return "Attachment metadata was not found."
    mime = str(metadata.get("mime") or "application/octet-stream").lower()
    if str(metadata.get("kind") or "") == "image" or mime.startswith("image/"):
        return "Image attachment. Analyze it explicitly with vlm_query using this attachment_id."
    storage_key = str(metadata.get("storage_key") or "")
    filename = str(metadata.get("filename") or "attachment")
    try:
        with tempfile.TemporaryDirectory(prefix="ouroboros-attachment-") as temp_dir:
            path = pathlib.Path(temp_dir) / pathlib.PurePosixPath(filename).name
            provider.copy_to_path(
                SCOPE_UPLOADS,
                storage_key,
                path,
                max_bytes=MAX_DOCUMENT_BYTES,
            )
            content = run_parser(path, mime, mode="extract")
    except FileNotFoundError:
        return "Attachment content was not found."
    except ObjectTooLargeError:
        return "Attachment is too large to parse."
    except AttachmentError as exc:
        return f"Failed to parse attachment: {exc.reason}"
    return f"Attachment: {filename}\nMIME: {mime}\n\n{content}"


def get_tools() -> list[ToolEntry]:
    return [
        ToolEntry(
            name="attachment_read",
            schema=ToolSchema(
                name="attachment_read",
                description="Read a TXT, MD, CSV, JSON, PDF, XLSX, or DOCX attachment bound to "
                "the current task. Images must be analyzed with vlm_query.",
                parameters={
                    "type": "object",
                    "properties": {
                        "attachment_id": {
                            "type": "string",
                            "description": "Attachment id from the current user message",
                        }
                    },
                    "required": ["attachment_id"],
                },
            ),
            handler=_attachment_read,
            timeout_sec=60,
        )
    ]


__all__ = [
    "ValidatedAttachment",
    "attachment_mime_for_suffix",
    "get_tools",
    "run_parser",
    "validate_attachment_path",
]

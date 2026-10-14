from __future__ import annotations

import pathlib

import pytest
from conftest import RunScript


def test_builds_document_with_heading_and_paragraph(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    code, payload = run_script("make_docx.py", tmp_path, "--out", "report.docx", "--text", "# Отчёт\n\nПервый абзац.")

    assert code == 0
    assert payload["status"] == "ok"
    assert payload["file"] == "report.docx"

    from docx import Document

    document = Document(str(tmp_path / "report.docx"))
    assert [p.text for p in document.paragraphs if p.text] == ["Отчёт", "Первый абзац."]
    assert document.paragraphs[0].style.name.startswith("Heading")


def test_renders_bullet_and_numbered_lists(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    code, payload = run_script(
        "make_docx.py", tmp_path, "--out", "lists.docx", "--text", "- первый\n- второй\n\n1. шаг\n2. шаг"
    )

    assert (code, payload["status"]) == (0, "ok")

    from docx import Document

    styles = [p.style.name for p in Document(str(tmp_path / "lists.docx")).paragraphs if p.text]
    assert styles == ["List Bullet", "List Bullet", "List Number", "List Number"]


def test_each_numbered_list_starts_at_one(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    # Word numbers every List Number paragraph of a document as one list unless a paragraph
    # points at its own numbering instance; without that the second list here would show 3 and 4.
    text = "1. первый\n2. второй\n\nАбзац между списками.\n\n1. снова первый\n2. снова второй"
    code, payload = run_script("make_docx.py", tmp_path, "--out", "numbered.docx", "--text", text)

    assert (code, payload["status"]) == (0, "ok")

    from docx import Document

    document = Document(str(tmp_path / "numbered.docx"))
    ids = [p._p.pPr.numPr.numId.val for p in document.paragraphs if p.style.name == "List Number"]
    assert ids[0] == ids[1] != ids[2] == ids[3]
    second = document.part.numbering_part.element.num_having_numId(ids[2])
    assert second.lvlOverride_lst[0].startOverride.val == 1


def test_renders_pipe_table_with_bold_header(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    text = "| Регион | Сумма |\n| --- | --- |\n| Москва | 10 |"
    code, payload = run_script("make_docx.py", tmp_path, "--out", "table.docx", "--text", text)

    assert (code, payload["status"]) == (0, "ok")

    from docx import Document

    table = Document(str(tmp_path / "table.docx")).tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Регион", "Сумма"]
    assert [cell.text for cell in table.rows[1].cells] == ["Москва", "10"]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True


def test_a_row_of_single_dashes_is_data_not_a_header_separator(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    text = "| a | b |\n| - | - |\n| c | d |"
    code, _ = run_script("make_docx.py", tmp_path, "--out", "dashes.docx", "--text", text)

    assert code == 0

    from docx import Document

    table = Document(str(tmp_path / "dashes.docx")).tables[0]
    assert [[cell.text for cell in row.cells] for row in table.rows] == [["a", "b"], ["-", "-"], ["c", "d"]]
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is None


def test_applies_inline_emphasis(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    code, _ = run_script("make_docx.py", tmp_path, "--out", "emphasis.docx", "--text", "обычный **жирный** и *курсив*")

    assert code == 0

    from docx import Document

    runs = Document(str(tmp_path / "emphasis.docx")).paragraphs[0].runs
    assert [(run.text, bool(run.bold), bool(run.italic)) for run in runs] == [
        ("обычный ", False, False),
        ("жирный", True, False),
        (" и ", False, False),
        ("курсив", False, True),
    ]


def test_a_backslash_keeps_markup_characters_literal(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    text = "1\\. Предмет договора\n\n2 \\* 3 \\* 4\n\n\\# не заголовок\n\nпуть C:\\\\dir и a \\| b"
    code, _ = run_script("make_docx.py", tmp_path, "--out", "escaped.docx", "--text", text)

    assert code == 0

    from docx import Document

    paragraphs = Document(str(tmp_path / "escaped.docx")).paragraphs
    assert [p.text for p in paragraphs] == [
        "1. Предмет договора",
        "2 * 3 * 4",
        "# не заголовок",
        "путь C:\\dir и a | b",
    ]
    assert all(p.style.name == "Normal" and not any(run.italic for run in p.runs) for p in paragraphs)


def test_joins_several_text_parts_in_order(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    code, _ = run_script(
        "make_docx.py", tmp_path, "--out", "parts.docx", "--text", "# Заголовок", "--text", "первая", "--text", "вторая"
    )

    assert code == 0

    from docx import Document

    assert [p.text for p in Document(str(tmp_path / "parts.docx")).paragraphs] == ["Заголовок", "первая вторая"]


@pytest.mark.parametrize("name", ["sub/escaped.docx", "../escaped.docx"], ids=["subdirectory", "parent directory"])
def test_refuses_a_destination_that_is_not_a_bare_file_name(
    run_script: RunScript, tmp_path: pathlib.Path, name: str
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()

    code, payload = run_script("make_docx.py", workspace, "--out", name, "--text", "текст")

    assert code == 2
    assert payload["status"] == "error"
    assert "not a bare file name" in payload["error"]
    assert not list(tmp_path.rglob("escaped.docx"))


def test_refuses_to_run_without_a_working_directory(
    run_script: RunScript, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Outside skill_exec there is no workdir; the current directory must not become the fallback.
    monkeypatch.chdir(tmp_path)

    code, payload = run_script("make_docx.py", None, "--out", "stray.docx", "--text", "текст")

    assert code == 2
    assert "OUROBOROS_SKILL_WORKDIR" in payload["error"]
    assert not (tmp_path / "stray.docx").exists()


@pytest.mark.parametrize("text_args", [(), ("--text", "   ")], ids=["no text", "blank text"])
def test_refuses_empty_input(run_script: RunScript, tmp_path: pathlib.Path, text_args: tuple[str, ...]) -> None:
    code, payload = run_script("make_docx.py", tmp_path, "--out", "empty.docx", *text_args)

    assert code == 2
    assert payload["status"] == "error"


def test_strips_control_characters_that_word_cannot_store(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    code, payload = run_script("make_docx.py", tmp_path, "--out", "control.docx", "--text", "стра\x0cница\x01текст")

    assert (code, payload["status"]) == (0, "ok")

    from docx import Document

    assert Document(str(tmp_path / "control.docx")).paragraphs[0].text == "страницатекст"

#!/usr/bin/env python3
"""Read a .docx back into the light markup make_docx.py accepts."""

from __future__ import annotations

import argparse
import sys

import common
import docx
from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def run_markup(runs: list[Run], *, plain_bold: bool = False) -> str:
    """Join runs into text, marking bold and italic the way make_docx.py reads them.

    A table header row is bold by construction, so there bold is the row marker, not emphasis.
    """
    merged: list[tuple[str, bool, bool]] = []
    for run in runs:
        if not run.text:
            continue
        bold, italic = bool(run.bold) and not plain_bold, bool(run.italic)
        if merged and merged[-1][1:] == (bold, italic):
            merged[-1] = (merged[-1][0] + run.text, bold, italic)
        else:
            merged.append((run.text, bold, italic))
    parts: list[str] = []
    for text, bold, italic in merged:
        # The markup has no bold-italic form, so bold wins when a run carries both.
        text = common.escape_text(text)
        if bold:
            parts.append(f"**{text}**")
        elif italic:
            parts.append(f"*{text}*")
        else:
            parts.append(text)
    return "".join(parts)


def paragraph_line(paragraph: Paragraph, ordinal: int) -> tuple[str, str, int]:
    """Return (kind, markup line, next ordered-list ordinal) for *paragraph*."""
    style = paragraph.style.name if paragraph.style is not None else ""
    text = run_markup(paragraph.runs)
    if style.startswith("Heading"):
        suffix = style.split()[-1]
        level = int(suffix) if suffix.isdigit() else 1
        return "block", f"{'#' * level} {text}", 0
    if style.startswith("List Number"):
        return "number", f"{ordinal + 1}. {text}", ordinal + 1
    if style.startswith("List Bullet"):
        return "bullet", f"- {text}", 0
    return "block", common.escape_line_start(text), 0


def cell_markup(paragraphs: list[Paragraph], *, header: bool) -> str:
    runs = [run for paragraph in paragraphs for run in paragraph.runs]
    return " ".join(run_markup(runs, plain_bold=header).split())


def table_lines(table: Table) -> list[str]:
    if not table.rows:
        return []
    first = table.rows[0].cells
    first_runs = [run for cell in first for paragraph in cell.paragraphs for run in paragraph.runs if run.text.strip()]
    header = bool(first_runs) and all(run.bold for run in first_runs)
    lines = [
        "| " + " | ".join(cell_markup(cell.paragraphs, header=header and index == 0) for cell in row.cells) + " |"
        for index, row in enumerate(table.rows)
    ]
    if header:
        lines.insert(1, "| " + " | ".join("---" for _ in first) + " |")
    return lines


def to_markup(document: Document) -> str:
    """Walk the body in document order; consecutive items of one list kind form one chunk."""
    chunks: list[str] = []
    run_kind = ""
    run_lines: list[str] = []
    ordinal = 0

    def flush() -> None:
        nonlocal run_kind
        if run_lines:
            chunks.append("\n".join(run_lines))
            run_lines.clear()
        run_kind = ""

    for block in document.iter_inner_content():
        if isinstance(block, Table):
            flush()
            ordinal = 0
            lines = table_lines(block)
            if lines:
                chunks.append("\n".join(lines))
            continue
        kind, line, ordinal = paragraph_line(block, ordinal)
        if not line.strip():
            flush()
            continue
        if kind == "block":
            flush()
            chunks.append(line)
            continue
        if kind != run_kind:
            flush()
            run_kind = kind
        run_lines.append(line)
    flush()
    return "\n\n".join(chunks)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Read a .docx back into light markup.")
    parser.add_argument("--in", dest="source", required=True, help="drive path of a .docx this skill built")
    args = parser.parse_args(argv)

    path = common.open_input(args.source)
    common.guard_package(path, "word/document.xml")
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx has no single error class for a malformed package
        common.fail(f"{path.name} is not a readable .docx: {exc}", status="invalid_input")

    text, truncated = common.bound_output(to_markup(document))
    common.emit(
        {
            "status": "ok",
            "file": path.name,
            "bytes": path.stat().st_size,
            "text": text,
            "truncated": truncated,
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

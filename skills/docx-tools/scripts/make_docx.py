#!/usr/bin/env python3
"""Build a .docx from plain or lightly structured text."""

from __future__ import annotations

import argparse
import sys

import common
import docx
from docx.document import Document
from docx.text.paragraph import Paragraph


def write_rich_text(paragraph: Paragraph, text: str) -> None:
    """Write *text* into *paragraph*, turning emphasis markers into styled runs."""
    for fragment, style in common.split_emphasis(text):
        run = paragraph.add_run(fragment)
        if style == "bold":
            run.bold = True
        elif style == "italic":
            run.italic = True


def add_table(document: Document, block: common.Table) -> None:
    width = max(len(row) for row in block.rows)
    table = document.add_table(rows=len(block.rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(block.rows):
        for column_index in range(width):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            write_rich_text(paragraph, row[column_index] if column_index < len(row) else "")
            if block.header and row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def new_numbering(document: Document) -> int:
    """Return a numId that starts at 1: the List Number style alone would continue the previous list."""
    numbering = document.part.numbering_part.element
    style_num_id = document.styles["List Number"].element.pPr.numPr.numId.val
    num = numbering.add_num(numbering.num_having_numId(style_num_id).abstractNumId.val)
    num.add_lvlOverride(0).add_startOverride(1)
    return num.numId


def render(document: Document, blocks: list[common.Block]) -> None:
    num_id: int | None = None
    for block in blocks:
        ordered_item = isinstance(block, common.ListItem) and block.ordered
        if not ordered_item:
            num_id = None
        if isinstance(block, common.Heading):
            write_rich_text(document.add_heading("", block.level), block.text)
        elif isinstance(block, common.ListItem):
            paragraph = document.add_paragraph(style="List Number" if block.ordered else "List Bullet")
            write_rich_text(paragraph, block.text)
            if ordered_item:
                num_id = new_numbering(document) if num_id is None else num_id
                # python-docx exposes no API for a paragraph's numbering instance, only the raw element.
                paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
        elif isinstance(block, common.Table):
            add_table(document, block)
        else:
            write_rich_text(document.add_paragraph(), block.text)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Build a .docx from text.")
    parser.add_argument("--out", required=True, help="file name of the .docx to build, without directories")
    parser.add_argument(
        "--text",
        action="append",
        default=[],
        help="document text; repeat for a long document, parts are joined in order",
    )
    args = parser.parse_args(argv)

    blocks = common.parse_blocks(common.clean_text(args.text))
    if not blocks:
        common.fail("nothing to write: the input text is empty")

    out_path = common.workdir_file(args.out, must_exist=False)
    if out_path.suffix.lower() != ".docx":
        common.fail(f"destination must end with .docx: {out_path.name}")

    document = docx.Document()
    render(document, blocks)
    try:
        document.save(str(out_path))
    except OSError as exc:
        common.fail(f"cannot write {out_path.name}: {exc}")

    common.emit(
        {
            "status": "ok",
            "file": out_path.name,
            "bytes": out_path.stat().st_size,
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

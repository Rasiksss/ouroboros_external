#!/usr/bin/env python3
"""Build an .xlsx workbook from JSON rows."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass

import common
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_SHEET_NAME_CHARS = 31
# Excel autofit is not stored in the file, so a width is set per column; past 60 characters a
# column stops fitting on screen and long text is better left to wrap over the neighbour.
MAX_COLUMN_WIDTH = 60
_INVALID_SHEET_CHARS = str.maketrans(dict.fromkeys("[]:*?/\\", "-"))
# The TSV form is what attachment_read prints for a spreadsheet: one "## Sheet:"
# line per sheet, cells separated by tabs.
_SHEET_HEADER_RE = re.compile(r"^## Sheet: (.*)$")
# Only plain decimals become numbers; "007", "+7900...", "1,5" and "1e5" stay text,
# because a code or a phone number must not be reshaped by Excel.
_NUMBER_RE = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


@dataclass(frozen=True)
class Sheet:
    name: str
    rows: list[list[object]]
    header: bool


def parse_sheets(raw: str) -> list[Sheet]:
    """Accept either a bare table or a {"sheets": [...]} document."""
    try:
        document = json.loads(raw)
    except json.JSONDecodeError as exc:
        common.fail(f"input is not valid JSON: {exc}", status="invalid_input")
    if isinstance(document, list):
        document = {"sheets": [{"name": "Sheet1", "rows": document}]}
    if not isinstance(document, dict) or not isinstance(document.get("sheets"), list):
        common.fail(
            'expected a list of rows or an object with a \'sheets\' list, e.g. [["Region","Total"],["Moscow",10]]',
            status="invalid_input",
        )
    sheets: list[Sheet] = []
    for index, sheet in enumerate(document["sheets"], start=1):
        if not isinstance(sheet, dict):
            common.fail(f"sheet #{index} is not an object", status="invalid_input")
        rows = sheet.get("rows")
        if not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
            common.fail(f"sheet #{index} must carry 'rows' as a list of lists", status="invalid_input")
        sheets.append(
            Sheet(
                name=sheet_name(sheet.get("name") or f"Sheet{index}"),
                rows=rows,
                header=bool(sheet.get("header")),
            )
        )
    if not sheets:
        common.fail("the workbook has no sheets", status="invalid_input")
    return sheets


def tsv_value(raw: str) -> common.CellValue:
    if raw == "":
        return None
    if _NUMBER_RE.match(raw):
        return float(raw) if "." in raw else int(raw)
    return raw


def parse_tsv(raw: str) -> list[Sheet]:
    """Accept the tab-separated dump attachment_read produces, with or without sheet headers."""
    named: list[tuple[str, list[list[object]]]] = []
    for line in raw.rstrip("\n").split("\n"):
        header = _SHEET_HEADER_RE.match(line)
        if header:
            named.append((header.group(1).strip(), []))
            continue
        if not named:
            named.append(("", []))
        cells: list[object] = [tsv_value(cell) for cell in line.rstrip("\r").split("\t")]
        # attachment_read pads every row to its column cap with empty cells; a
        # trailing empty cell and no cell are the same thing in a spreadsheet.
        while cells and cells[-1] is None:
            cells.pop()
        named[-1][1].append(cells)
    return [
        Sheet(name=sheet_name(name or f"Sheet{index}"), rows=rows, header=False)
        for index, (name, rows) in enumerate(named, start=1)
    ]


def sheet_name(raw: object) -> str:
    """Return a name Excel accepts: no reserved characters, at most 31 chars."""
    name = str(raw).translate(_INVALID_SHEET_CHARS).strip() or "Sheet"
    return name[:MAX_SHEET_NAME_CHARS]


def write_cell(sheet: Worksheet, row_index: int, column_index: int, value: object) -> Cell:
    cell = sheet.cell(row=row_index, column=column_index)
    if isinstance(value, (int, float, bool)) or value is None:
        cell.value = value
        return cell
    cell.value = common.strip_control_characters(value if isinstance(value, str) else str(value))
    # openpyxl types a leading "=" as a formula, which Excel would then execute;
    # user-supplied text must stay text.
    cell.data_type = "s"
    return cell


def fit_columns(sheet: Worksheet, rows: list[list[object]]) -> None:
    """Widen each column to its longest value so a header or a name is not clipped by the neighbour."""
    longest: dict[int, int] = {}
    for row in rows:
        for column_index, value in enumerate(row, start=1):
            if value is None:
                continue
            width = max(len(line) for line in str(value).split("\n"))
            longest[column_index] = max(longest.get(column_index, 0), width)
    for column_index, width in longest.items():
        sheet.column_dimensions[get_column_letter(column_index)].width = min(MAX_COLUMN_WIDTH, width + 2)


def build(workbook: Workbook, sheets: list[Sheet]) -> None:
    workbook.remove(workbook.active)
    for spec in sheets:
        sheet = workbook.create_sheet(title=spec.name)
        for row_index, row in enumerate(spec.rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = write_cell(sheet, row_index, column_index, value)
                if spec.header and row_index == 1:
                    cell.font = Font(bold=True)
        fit_columns(sheet, spec.rows)


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Build an .xlsx workbook from JSON or TSV rows.")
    parser.add_argument("--out", required=True, help="file name of the .xlsx to build, without directories")
    parser.add_argument(
        "--text",
        action="append",
        default=[],
        help="JSON or TSV source; repeat for a large table, parts are joined in order",
    )
    parser.add_argument("--format", choices=["json", "tsv"], default="json", help="source format")
    args = parser.parse_args(argv)

    # JSON tolerates a split anywhere between two tokens, TSV only between two lines.
    raw = ("\n" if args.format == "tsv" else "").join(args.text)
    if not raw.strip():
        common.fail("nothing to write: the input is empty")
    sheets = parse_tsv(raw) if args.format == "tsv" else parse_sheets(raw)

    out_path = common.workdir_file(args.out, must_exist=False)
    if out_path.suffix.lower() != ".xlsx":
        common.fail(f"destination must end with .xlsx: {out_path.name}")

    workbook = Workbook()
    build(workbook, sheets)
    try:
        workbook.save(str(out_path))
    except OSError as exc:
        common.fail(f"cannot write {out_path.name}: {exc}")

    common.emit(
        {
            "status": "ok",
            "file": out_path.name,
            "bytes": out_path.stat().st_size,
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

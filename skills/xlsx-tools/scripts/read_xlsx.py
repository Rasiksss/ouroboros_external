#!/usr/bin/env python3
"""Read an .xlsx workbook into the JSON shape make_xlsx.py accepts."""

from __future__ import annotations

import argparse
import datetime
import json
import sys
from contextlib import closing
from dataclasses import asdict, dataclass, field

import common
from openpyxl import Workbook, load_workbook


@dataclass(frozen=True)
class SheetRows:
    name: str
    rows: list[list[common.CellValue]] = field(default_factory=list)


def cell_value(value: object) -> common.CellValue:
    """Return *value* as something JSON can carry and make_xlsx.py can write back."""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (bool, int, float, str)) or value is None:
        return value
    return str(value)


def read_sheets(workbook: Workbook) -> tuple[list[SheetRows], bool]:
    """Collect rows per sheet until the output budget is spent; report whether it was."""
    sheets: list[SheetRows] = []
    length = 0
    for sheet in workbook.worksheets:
        current = SheetRows(name=sheet.title)
        sheets.append(current)
        for row in sheet.iter_rows(values_only=True):
            values = [cell_value(value) for value in row]
            length += len(json.dumps(values, ensure_ascii=False))
            if length > common.MAX_OUTPUT_CHARS:
                return sheets, True
            current.rows.append(values)
    return sheets, False


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Read an .xlsx workbook into JSON rows.")
    parser.add_argument("--in", dest="source", required=True, help="drive path of an .xlsx this skill built")
    args = parser.parse_args(argv)

    path = common.open_input(args.source)
    common.guard_package(path, "xl/workbook.xml")
    try:
        # data_only returns the value Excel last cached for a formula, never the formula;
        # read_only streams the sheets, so a malformed one surfaces here, not at open time.
        with closing(load_workbook(path, read_only=True, data_only=True)) as workbook:
            sheets, truncated = read_sheets(workbook)
    except Exception as exc:  # openpyxl has no single error class for a malformed package
        common.fail(f"{path.name} is not a readable .xlsx: {exc}", status="invalid_input")

    common.emit(
        {
            "status": "ok",
            "file": path.name,
            "bytes": path.stat().st_size,
            "sheets": [asdict(sheet) for sheet in sheets],
            "truncated": truncated,
        }
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))

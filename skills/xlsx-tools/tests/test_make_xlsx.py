from __future__ import annotations

import json
import pathlib

import pytest
from conftest import RunScript


def test_writes_rows_preserving_value_types(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "data.xlsx"
    rows = json.dumps([["Регион", "Выручка"], ["Москва", 1250.5], ["Казань", 300]])

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--text", rows)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    sheet = load_workbook(out)["Sheet1"]
    assert [cell.value for cell in sheet[1]] == ["Регион", "Выручка"]
    assert [cell.value for cell in sheet[2]] == ["Москва", 1250.5]
    assert sheet["B3"].value == 300


def test_writes_named_sheets_with_bold_header(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "book.xlsx"
    payload_in = json.dumps(
        {
            "sheets": [
                {"name": "Продажи", "header": True, "rows": [["Месяц", "Сумма"], ["Январь", 10]]},
                {"name": "Затраты", "rows": [["Статья", "Сумма"]]},
            ]
        }
    )

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--text", payload_in)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    book = load_workbook(out)
    assert book.sheetnames == ["Продажи", "Затраты"]
    assert book["Продажи"]["A1"].font.bold is True
    assert book["Затраты"]["A1"].font.bold is not True


def test_columns_are_widened_to_their_longest_value_up_to_a_cap(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "wide.xlsx"
    rows = json.dumps([["Регион", "Выручка, млн руб.", "x" * 200], ["Санкт-Петербург", 63.7, None]])

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--text", rows)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    columns = load_workbook(out)["Sheet1"].column_dimensions
    assert columns["A"].width == len("Санкт-Петербург") + 2
    assert columns["B"].width == len("Выручка, млн руб.") + 2
    assert columns["C"].width == 60


def test_never_writes_user_text_as_a_formula(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "injection.xlsx"
    rows = json.dumps([["=SUM(B1:B9)", "+79001234567", "@user"]])

    code, _ = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--text", rows)

    assert code == 0

    from openpyxl import load_workbook

    sheet = load_workbook(out).active
    assert [cell.data_type for cell in sheet[1]] == ["s", "s", "s"]
    assert sheet["A1"].value == "=SUM(B1:B9)"


@pytest.mark.parametrize(
    "source",
    ["{not json", json.dumps({"rows": "нет"}), json.dumps({"sheets": []})],
    ids=["broken-json", "rows-not-a-table", "no-sheets"],
)
def test_refuses_input_that_is_not_a_table(run_script: RunScript, tmp_path: pathlib.Path, source: str) -> None:
    code, payload = run_script("make_xlsx.py", tmp_path, "--out", "rejected.xlsx", "--text", source)

    assert code == 2
    assert payload["status"] == "invalid_input"


@pytest.mark.parametrize("text_args", [(), ("--text", "  ")], ids=["no text", "blank text"])
def test_refuses_empty_input(run_script: RunScript, tmp_path: pathlib.Path, text_args: tuple[str, ...]) -> None:
    code, payload = run_script("make_xlsx.py", tmp_path, "--out", "empty.xlsx", *text_args)

    assert code == 2
    assert payload["status"] == "error"


@pytest.mark.parametrize(
    ("format_args", "parts", "expected"),
    [
        ((), ('[["a", "b"], ["c"', ', "d"]]'), [["a", "b"], ["c", "d"]]),
        (("--format", "tsv"), ("a\tb", "c\td"), [["a", "b"], ["c", "d"]]),
    ],
    ids=["json split inside a row", "tsv split between rows"],
)
def test_joins_several_text_parts_in_order(
    run_script: RunScript, tmp_path: pathlib.Path, format_args: tuple[str, ...], parts: tuple[str, ...], expected: list
) -> None:
    text_args = [arg for part in parts for arg in ("--text", part)]
    code, payload = run_script("make_xlsx.py", tmp_path, "--out", "parts.xlsx", *format_args, *text_args)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    assert [[cell.value for cell in row] for row in load_workbook(tmp_path / "parts.xlsx").active.rows] == expected


@pytest.mark.parametrize("name", ["sub/escaped.xlsx", "../escaped.xlsx"], ids=["subdirectory", "parent directory"])
def test_refuses_a_destination_that_is_not_a_bare_file_name(
    run_script: RunScript, tmp_path: pathlib.Path, name: str
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()

    code, payload = run_script("make_xlsx.py", workspace, "--out", name, "--text", json.dumps([["a"]]))

    assert code == 2
    assert "not a bare file name" in payload["error"]
    assert not list(tmp_path.rglob("escaped.xlsx"))


def test_refuses_to_run_without_a_working_directory(
    run_script: RunScript, tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Outside skill_exec there is no workdir; the current directory must not become the fallback.
    monkeypatch.chdir(tmp_path)

    code, payload = run_script("make_xlsx.py", None, "--out", "stray.xlsx", "--text", json.dumps([["a"]]))

    assert code == 2
    assert "OUROBOROS_SKILL_WORKDIR" in payload["error"]
    assert not (tmp_path / "stray.xlsx").exists()


def test_strips_control_characters_that_excel_rejects(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "control.xlsx"

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--text", json.dumps([["зна\x00чение"]]))

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    assert load_workbook(out).active["A1"].value == "значение"


def test_accepts_the_tsv_attachment_read_produces(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "from-tsv.xlsx"
    tsv = "## Sheet: Продажи\nМесяц\tСумма\tТелефон\tКод\tПусто\nЯнварь\t10\t+79001234567\t007\t\n## Sheet: Затраты\nСтатья\tСумма\nАренда\t1.5"

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--format", "tsv", "--text", tsv)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    book = load_workbook(out)
    assert book.sheetnames == ["Продажи", "Затраты"]
    assert [cell.value for cell in book["Продажи"][2]] == ["Январь", 10, "+79001234567", "007", None]
    assert book["Затраты"]["B2"].value == 1.5


def test_tsv_without_sheet_headers_becomes_one_sheet(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    out = tmp_path / "plain-tsv.xlsx"

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--format", "tsv", "--text", "a\tb\n1\t2\n")

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    assert load_workbook(out).sheetnames == ["Sheet1"]


def test_tsv_drops_the_empty_columns_attachment_read_pads_rows_with(
    run_script: RunScript, tmp_path: pathlib.Path
) -> None:
    out = tmp_path / "padded.xlsx"
    padded = "## Sheet: S\nМесяц\tСумма" + "\t" * 98 + "\nЯнварь\t10" + "\t" * 98

    code, payload = run_script("make_xlsx.py", tmp_path, "--out", out.name, "--format", "tsv", "--text", padded)

    assert (code, payload["status"]) == (0, "ok")

    from openpyxl import load_workbook

    assert load_workbook(out).active.max_column == 2

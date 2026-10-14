from __future__ import annotations

import datetime
import json
import pathlib
import zipfile

import pytest
from conftest import RunScript

WORKBOOK = {
    "sheets": [{"name": "Продажи", "header": True, "rows": [["Месяц", "Сумма"], ["Январь", 10], ["Февраль", 1.5]]}]
}


def build(run_script: RunScript, workspace: pathlib.Path, document: dict, name: str = "book.xlsx") -> pathlib.Path:
    out = workspace / name
    code, payload = run_script("make_xlsx.py", workspace, "--out", out.name, "--text", json.dumps(document))
    assert (code, payload["status"]) == (0, "ok")
    return out


def test_reads_sheets_in_the_shape_make_xlsx_accepts(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    built = build(run_script, tmp_path, WORKBOOK)

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", built.name)

    assert (code, payload["status"]) == (0, "ok")
    assert payload["sheets"] == [{"name": "Продажи", "rows": [["Месяц", "Сумма"], ["Январь", 10], ["Февраль", 1.5]]}]
    assert payload["truncated"] is False


@pytest.mark.parametrize(
    "name", ["missing.xlsx", "uploads/missing.xlsx"], ids=["bare name", "drive path passed through"]
)
def test_refuses_an_input_that_is_not_among_the_call_inputs(
    run_script: RunScript, tmp_path: pathlib.Path, name: str
) -> None:
    code, payload = run_script("read_xlsx.py", tmp_path, "--in", name)

    assert code == 2
    assert payload["status"] == "error"
    assert payload["error"].startswith(f"{name} is not among this call's inputs")


def test_an_empty_sheet_survives_the_round_trip(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Данные"
    book.active.append(["a", 1])
    book.create_sheet("Пусто")
    source = tmp_path / "with-empty.xlsx"
    book.save(source)

    _, first_read = run_script("read_xlsx.py", tmp_path, "--in", source.name)
    rebuilt = build(run_script, tmp_path, {"sheets": first_read["sheets"]}, "rebuilt.xlsx")
    _, second_read = run_script("read_xlsx.py", tmp_path, "--in", rebuilt.name)

    assert first_read["sheets"] == [{"name": "Данные", "rows": [["a", 1]]}, {"name": "Пусто", "rows": []}]
    assert second_read["sheets"] == first_read["sheets"]


def test_dates_come_back_as_iso_strings_and_formulas_as_cached_values(
    run_script: RunScript, tmp_path: pathlib.Path
) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet["A1"] = datetime.datetime(2026, 8, 18, 9, 30)
    sheet["B1"] = "=1+1"
    path = tmp_path / "typed.xlsx"
    book.save(path)

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", path.name)

    assert (code, payload["status"]) == (0, "ok")
    # A workbook never opened in Excel carries no cached result for its formulas.
    assert payload["sheets"][0]["rows"] == [["2026-08-18T09:30:00", None]]


def write_package(path: pathlib.Path, members: dict[str, str | bytes]) -> pathlib.Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in members.items():
            archive.writestr(name, body)
    return path


def broken_sheet(run_script: RunScript, workspace: pathlib.Path) -> dict[str, str | bytes]:
    """A workbook make_xlsx built, with its only sheet part replaced by garbage."""
    with zipfile.ZipFile(build(run_script, workspace, WORKBOOK)) as archive:
        members: dict[str, str | bytes] = {name: archive.read(name) for name in archive.namelist()}
    members["xl/worksheets/sheet1.xml"] = "not xml <<<"
    return members


@pytest.mark.parametrize(
    "case",
    ["not a zip", "workbook part alone", "sheet part is not xml"],
)
def test_refuses_a_package_openpyxl_cannot_read(run_script: RunScript, tmp_path: pathlib.Path, case: str) -> None:
    fake = tmp_path / "fake.xlsx"
    if case == "not a zip":
        fake.write_text("not a package", encoding="utf-8")
    elif case == "workbook part alone":
        write_package(fake, {"xl/workbook.xml": "<workbook/>"})
    else:
        write_package(fake, broken_sheet(run_script, tmp_path))

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", fake.name)

    assert code == 2
    assert payload["status"] == "invalid_input"


def test_refuses_a_package_that_unpacks_far_beyond_its_size(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    # A few KB on disk, megabytes once inflated: the size check on the file alone would let it through.
    bomb = write_package(tmp_path / "bomb.xlsx", {"xl/workbook.xml": "<workbook>" + " " * 5_000_000})

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", bomb.name)

    assert code == 2
    assert payload["status"] == "input_too_large"


def test_refuses_input_over_the_size_limit_before_opening_it(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    huge = tmp_path / "huge.xlsx"
    with huge.open("wb") as handle:
        handle.truncate(50 * 1024 * 1024 + 1)

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", huge.name)

    assert code == 2
    assert payload["status"] == "input_too_large"


def test_stops_reading_at_the_output_limit_and_says_so(run_script: RunScript, tmp_path: pathlib.Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    # Varied cells, not one repeated string: a sheet that inflates 200:1 is refused as a package bomb.
    for row in range(1000):
        book.active.append([f"ячейка {row}-{column}, значение {row * column} из выборки" for column in range(4)])
    built = tmp_path / "big.xlsx"
    book.save(built)

    code, payload = run_script("read_xlsx.py", tmp_path, "--in", built.name)

    assert (code, payload["status"]) == (0, "ok")
    assert payload["truncated"] is True
    assert 0 < len(payload["sheets"][0]["rows"]) < 1000
    assert sum(len(json.dumps(row, ensure_ascii=False)) for row in payload["sheets"][0]["rows"]) <= 100_000
